#!/usr/bin/env python3
"""Build the manuscript, supplement, checklist, and response package for R2.

The script intentionally starts from the R1 Word files so that journal formatting,
references, and unchanged figures are retained. It writes new R2 files and never
overwrites the R1 source documents.
"""
from __future__ import annotations

import csv
import shutil
import zipfile
from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parent.parent
R1 = ROOT / ".cursor" / "low_fe_3"
R2 = ROOT / "deliverables" / "Fe_R2_Revision_Package"
REF = ROOT / "reference" / "r2"
REVIEW_SOURCES = R1 / "r2_review_sources"
UPDATED_FIGURE1 = REF / "figure1_biorender_updated.png"

NEW_TITLE = (
    "Mortality Risk Stratification Among Critically Ill Patients with Documented "
    "Hypoferremia: Multi-Cohort Development and Retrospective Evaluation"
)


def _set_document_metadata(doc: Document, title: str, subject: str) -> None:
    props = doc.core_properties
    props.title = title
    props.subject = subject
    props.keywords = "hypoferremia; intensive care; mortality prediction; machine learning"
    props.comments = "Round 2 revision"


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _remove_cell_shading(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is not None:
        tc_pr.remove(shd)


def _set_cell_border(cell, edge: str, *, value: str, size: int = 0) -> None:
    """Set one explicit cell border so table styles cannot reintroduce grid lines."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    node = borders.find(qn(f"w:{edge}"))
    if node is None:
        node = OxmlElement(f"w:{edge}")
        borders.append(node)
    node.set(qn("w:val"), value)
    if value == "single":
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "000000")


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def _apply_three_line_table(table, *, font_size: float = 8.0) -> None:
    """Apply the journal-style three-line table convention.

    Only the top rule, header-bottom rule, and bottom rule are retained. Cell
    shading, vertical rules, and internal body rules are removed explicitly so
    the result is stable across Word and LibreOffice renderers.
    """
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    tbl_pr = table._tbl.tblPr
    style = tbl_pr.find(qn("w:tblStyle"))
    if style is not None:
        tbl_pr.remove(style)
    look = tbl_pr.find(qn("w:tblLook"))
    if look is not None:
        look.set(qn("w:firstRow"), "0")
        look.set(qn("w:lastRow"), "0")
        look.set(qn("w:firstColumn"), "0")
        look.set(qn("w:lastColumn"), "0")
        look.set(qn("w:noHBand"), "1")
        look.set(qn("w:noVBand"), "1")

    if not table.rows:
        return
    _set_repeat_table_header(table.rows[0])
    last_row_index = len(table.rows) - 1
    for r_i, row in enumerate(table.rows):
        _keep_row_together(row)
        for c_i, cell in enumerate(row.cells):
            _remove_cell_shading(cell)
            _set_cell_margins(cell, top=55, start=60, bottom=55, end=60)
            for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
                _set_cell_border(cell, edge, value="nil")
            if r_i == 0:
                _set_cell_border(cell, "top", value="single", size=8)
                _set_cell_border(cell, "bottom", value="single", size=6)
            if r_i == last_row_index:
                _set_cell_border(cell, "bottom", value="single", size=8)

            text = cell.text.strip()
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER
                    if r_i == 0 or (c_i > 0 and len(text) <= 28)
                    else WD_ALIGN_PARAGRAPH.LEFT
                )
                for run in paragraph.runs:
                    _set_times_new_roman(run, size_pt=font_size)
                    run.bold = r_i == 0


def _set_cell_margins(cell, top=70, start=70, bottom=70, end=70) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def _keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def _set_times_new_roman(run, *, size_pt: float | None = None) -> None:
    """Apply the manuscript's R1 typeface explicitly to a run.

    The R1 files use Times New Roman through a mixture of style inheritance and
    direct run formatting. Explicitly setting all four Word font slots prevents
    newly inserted or highlighted text from falling back to a theme font.
    """
    run.font.name = "Times New Roman"
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{slot}"), "Times New Roman")
    if size_pt is not None:
        run.font.size = Pt(size_pt)


def _copy_paragraph_properties(source, target) -> None:
    """Copy direct paragraph geometry without changing the target text."""
    source_ppr = source._p.pPr
    target_ppr = target._p.pPr
    if target_ppr is not None:
        target._p.remove(target_ppr)
    if source_ppr is not None:
        target._p.insert(0, deepcopy(source_ppr))


def _template_font_size(paragraph, template_run, fallback: float = 12.0) -> float:
    if template_run is not None and template_run.font.size is not None:
        return template_run.font.size.pt
    style_size = paragraph.style.font.size
    if style_size is not None:
        return style_size.pt
    return fallback


def set_paragraph(
    paragraph,
    text: str,
    *,
    highlight: bool = False,
    bold: bool | None = None,
    font_size: float | None = None,
) -> None:
    """Replace paragraph text while preserving the R1 run hierarchy.

    Most R1 headings are encoded with direct run formatting rather than Word
    heading styles. The previous builder discarded that formatting and created
    theme-dependent replacement runs. Clone the first textual R1 run before
    clearing the paragraph, then apply only the requested revision attributes.
    """
    template_run = next((run for run in paragraph.runs if run.text), None)
    template_rpr = (
        deepcopy(template_run._r.rPr)
        if template_run is not None and template_run._r.rPr is not None
        else None
    )
    resolved_size = (
        font_size
        if font_size is not None
        else _template_font_size(paragraph, template_run)
    )
    paragraph.clear()
    run = paragraph.add_run(text)
    if template_rpr is not None:
        current_rpr = run._r.rPr
        if current_rpr is not None:
            run._r.remove(current_rpr)
        run._r.insert(0, template_rpr)
    _set_times_new_roman(run, size_pt=resolved_size)
    if bold is not None:
        run.bold = bold
    if highlight:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def set_text_preserving_drawing(
    paragraph,
    text: str,
    *,
    highlight: bool = False,
) -> None:
    """Replace a caption in a paragraph that also contains an inline figure."""
    drawing_runs = [run for run in paragraph.runs if run._r.xpath(".//w:drawing")]
    text_runs = [run for run in paragraph.runs if run.text]
    if not drawing_runs or not text_runs:
        raise ValueError("Expected a paragraph containing both an inline drawing and caption")
    target = text_runs[0]
    target.text = text
    _set_times_new_roman(
        target,
        size_pt=target.font.size.pt if target.font.size is not None else 12.0,
    )
    if highlight:
        target.font.highlight_color = WD_COLOR_INDEX.YELLOW
    for extra in text_runs[1:]:
        extra.text = ""


def replace_table(doc: Document, index: int, rows: list[list[str]], *, highlight: bool = False) -> None:
    old = doc.tables[index]
    new = doc.add_table(rows=len(rows), cols=len(rows[0]))
    new.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_i, row in enumerate(rows):
        _keep_row_together(new.rows[r_i])
        for c_i, value in enumerate(row):
            cell = new.cell(r_i, c_i)
            cell.text = "" if value is None else str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8)
                    run.bold = r_i == 0
                    if highlight:
                        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    _apply_three_line_table(new)
    old._element.addnext(new._element)
    old._element.getparent().remove(old._element)


def append_table(
    doc: Document,
    title: str,
    rows: list[list[str]],
    note: str | None = None,
    *,
    highlight: bool = False,
    caption_template=None,
    note_template=None,
) -> None:
    spacer = doc.add_paragraph()
    p = doc.add_paragraph()
    if caption_template is not None:
        _copy_paragraph_properties(caption_template, p)
    p.paragraph_format.keep_with_next = True
    set_paragraph(p, title, highlight=highlight, bold=True, font_size=12.0)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_i, row in enumerate(rows):
        _keep_row_together(table.rows[r_i])
        for c_i, value in enumerate(row):
            cell = table.cell(r_i, c_i)
            cell.text = "" if value is None else str(value)
            _set_cell_margins(cell)
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(8)
                    run.bold = r_i == 0
                    if highlight:
                        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    _apply_three_line_table(table)
    if note:
        p = doc.add_paragraph()
        if note_template is not None:
            _copy_paragraph_properties(note_template, p)
        set_paragraph(p, note, highlight=highlight, bold=False, font_size=10.0)


def df_rows(df: pd.DataFrame, *, digits: int = 3) -> list[list[str]]:
    def fmt(v):
        if pd.isna(v):
            return ""
        if isinstance(v, float):
            if abs(v) < 0.001 and v != 0:
                return f"{v:.2e}"
            return f"{v:.{digits}f}"
        return str(v)

    return [[str(x) for x in df.columns]] + [[fmt(v) for v in row] for row in df.itertuples(index=False, name=None)]


def load_table1() -> list[list[str]]:
    with (REF / "r2_table1_mimic4_eligible.csv").open(newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    # The export's second row contains the final display headers.
    return [rows[1]] + rows[2:]


MANUSCRIPT_REPLACEMENTS = {
    0: NEW_TITLE,
    13: (
        "Background: Low circulating serum iron is common during critical illness, but it may "
        "reflect inflammation-associated iron restriction rather than depleted iron stores. "
        "We evaluated mortality risk among ICU patients selected by documented biochemical "
        "hypoferremia; the analysis was predictive and does not establish a causal effect of iron."
    ),
    14: (
        "Methods: Adults with ICU length of stay ≥48 hours and serum iron <60 µg/dL were identified "
        "in MIMIC-IV. Of 1,487 eligible admissions, 1,344 had a complete locked 18-predictor matrix "
        "and were split 80:20 for model development and internal testing. The same locked models "
        "were retrospectively evaluated in a corrected MIMIC-III temporal cohort (n=799) and, for "
        "the post hoc 8-feature model, two exploratory Chinese cohorts (n=101 and n=98). Predictor "
        "minimum, maximum, and mean values summarized the full ICU stay. The endpoint was 28-day "
        "in-hospital mortality from hospital admission. Added analyses assessed calibration, "
        "multiple imputation, exclusion of the NIV documentation variable, and the adjusted and "
        "incremental contribution of first serum iron."
    ),
    15: (
        "Results: The locked 18-predictor soft-voting model achieved AUROC 0.840 (95% CI "
        "0.767–0.913) in the MIMIC-IV internal test set and 0.705 in the corrected MIMIC-III "
        "cohort. Internal calibration-in-the-large was −0.898, calibration slope 1.705, and Brier "
        "score 0.109. Five-imputation out-of-fold analysis of all 1,487 eligible MIMIC-IV admissions "
        "gave AUROC 0.808. First serum iron was not independently associated with mortality after "
        "adjustment (OR per 10 µg/dL, 1.05; 95% CI 0.95–1.17; P=0.348), showed no evidence of "
        "nonlinearity (P=0.555), and added no measurable discrimination (ΔAUROC −0.0014; 95% CI "
        "−0.0075 to 0.0044)."
    ),
    16: (
        "Conclusions: The models stratified retrospective mortality risk among patients selected "
        "for documented hypoferremia, but their signal was dominated by general severity markers "
        "rather than serum iron. Temporal performance loss and calibration error require local "
        "recalibration and prospective validation before any clinical use."
    ),
    17: (
        "Keywords: Hypoferremia; Intensive care unit; Mortality prediction; Machine learning; "
        "Temporal evaluation; Explainable AI."
    ),
    23: (
        "Here, we develop and retrospectively evaluate interpretable mortality-risk models among "
        "ICU patients selected by documented biochemical hypoferremia. The cohort definition "
        "identifies the population of interest; it does not imply that serum iron is an independent "
        "cause of mortality. We compare a locked 18-predictor benchmark with a post hoc 8-feature "
        "derivative and quantify discrimination, calibration, robustness to missing data, and the "
        "incremental contribution of serum iron."
    ),
    28: (
        "Model development and internal evaluation used MIMIC-IV v2.2. The available analysis "
        "export contained 1,931 ICU admissions before the ICU-length-of-stay filter. After requiring "
        "ICU length of stay ≥48 hours, 1,568 remained; 1,487 of these met the fixed serum-iron "
        "criterion (<60 µg/dL). The locked complete-feature analysis included 1,344 admissions. "
        "MIMIC-IV data are de-identified and were accessed through PhysioNet under its data-use requirements."
    ),
    29: "Non-training evaluation cohorts",
    30: (
        "Transportability was examined in one non-training same-institution temporal cohort and two "
        "small exploratory geographic cohorts:"
    ),
    31: (
        "MIMIC-III: a same-institution temporal cohort from an earlier era. Reapplying the stated "
        "ICU length-of-stay ≥48-hour criterion yielded 799 admissions with 140 endpoint events. "
        "This cohort was not used for feature selection, imputer fitting, or model training."
    ),
    32: (
        "Chinese exploratory geographic cohorts: consecutive adult admissions from the Third "
        "Xiangya Hospital (F3; n=101, 26 events) and the First Affiliated Hospital of the University "
        "of South China (NH; n=98, 37 events). These small cohorts evaluated only the 8-feature "
        "derivative and are not definitive independent validations."
    ),
    41: (
        "The final flow was: MIMIC-IV, 1,568 admissions after the ≥48-hour screen, 1,487 meeting "
        "serum iron <60 µg/dL, and 1,344 with a complete locked 18-predictor matrix; corrected "
        "MIMIC-III, n=799; F3, n=101; and NH, n=98. The 143 eligible MIMIC-IV admissions excluded "
        "from the complete matrix had 15 endpoint events. The sequential eligibility and "
        "complete-matrix denominators are reported explicitly in the text and Supplementary Table S7."
    ),
    42: (
        "A normal-serum-iron comparator analysis from the original submission could not be reproduced "
        "from the retained screening export because individual records for a valid comparator cohort "
        "and their extraction code were not available. That analysis and its Kaplan–Meier claim were "
        "removed. Figure 2 now reports only prespecified-compatible analyses within the 1,487 eligible "
        "hypoferremia admissions."
    ),
    44: "2a. Predictor observation window and analysis time origin",
    45: (
        "The retained analysis matrix did not contain authentic iron chart times and therefore does "
        "not support a first-low-iron prediction index. Variables ending in _min, _max, and _mean "
        "were calculated across the full ICU stay; variables ending in _1st were the first recorded "
        "value during the ICU stay. The existing los_icu_to_t0 field was a synthetic legacy proxy "
        "derived from length of stay and was not treated as a real timestamp. Accordingly, the model "
        "is presented as retrospective whole-stay risk stratification, not admission-time or real-time prediction."
    ),
    46: "3. Variables, preprocessing, and outcome definition",
    53: (
        "The primary endpoint was death in hospital within 28 days of hospital admission. The retained "
        "export provided hosp_survival_days measured from hospital admission; it did not contain a "
        "validated time-to-event origin at the first low-iron measurement. We therefore removed all "
        "statements that anchored follow-up to t0."
    ),
    54: (
        "Admissions discharged alive before day 28 were coded as non-events in the available in-hospital "
        "endpoint. This may introduce informative discharge bias because post-discharge deaths were not "
        "captured. Time-to-event analyses were not performed because comparable censoring and event times "
        "were unavailable across cohorts."
    ),
    55: (
        "4. Feature selection and data preprocessing. The primary locked analysis used 1,344 MIMIC-IV "
        "admissions with a complete 18-predictor matrix. A five-imputation, five-fold out-of-fold "
        "sensitivity analysis retained all 1,487 eligible admissions. External datasets were transformed "
        "using the training-fitted preprocessing objects. The NIV indicator denoted documented use, not "
        "physiological status; 468/799 (58.6%) values were missing in corrected MIMIC-III. We therefore "
        "also refitted models after excluding this variable."
    ),
    56: (
        "Predictor reduction used least-absolute-shrinkage-and-selection-operator (LASSO) logistic "
        "regression. The penalty was tuned by stratified ten-fold cross-validation using only the "
        "MIMIC-IV training data. The selected feature set was then fixed; the internal test set and "
        "all non-training evaluation cohorts were excluded from feature selection."
    ),
    58: (
        "5. Machine learning model development and evaluation. Analyses used Python and scikit-learn. "
        "Algorithms supporting class or sample weighting used balanced weights; other models were "
        "trained on the full imbalanced MIMIC-IV training set without random undersampling. Calibration "
        "and decision-curve analyses retained each evaluation cohort's native event prevalence. A "
        "five-seed random-undersampling sensitivity analysis is reported in Supplementary Table S8."
    ),
    60: (
        "The 18-predictor soft-voting model was retained as the prespecified benchmark, not because it "
        "outperformed simpler models. Logistic regression and all other algorithms were reported as "
        "comparators. Performance assessment included AUROC with confidence intervals, Brier score, "
        "calibration-in-the-large (intercept with the prediction logit as an offset), calibration slope, "
        "and retrospective decision-curve analysis. Calibration estimates from small samples were "
        "interpreted cautiously."
    ),
    61: (
        "SHAP analyses described predictive contributions and were not interpreted causally. The "
        "8-feature model was a secondary, post hoc implementation-oriented derivative based on the "
        "development data; it was not a replacement for the benchmark and its internal performance "
        "may be optimistic. Reporting was audited against TRIPOD+AI. The repository supplies the "
        "prepared-table analysis workflow, model objects, and aggregate outputs; it does not claim to "
        "recreate the protected modeling tables from raw clinical databases."
    ),
    62: "6. Additional serum-iron analyses",
    63: (
        "The association analysis included all 1,487 eligible MIMIC-IV admissions. We fitted logistic "
        "regression for first recorded serum iron per 10 µg/dL, adjusting for age, sex, race as an "
        "unordered categorical variable (White reference), SOFA score, Charlson comorbidity index, "
        "first lactate, invasive ventilation, and vasoactive therapy. Numeric covariates were median "
        "imputed and categorical covariates were mode imputed for the primary explanatory model; a "
        "complete-case sensitivity analysis used the same specification. Nonlinearity was assessed "
        "with a four-knot restricted cubic spline at the 5th, 35th, 65th, and 95th percentiles "
        "(9, 20, 33, and 55 µg/dL) and a likelihood-ratio test against the linear model. Incremental "
        "predictive value was evaluated with identical stratified five-fold splits for models with and "
        "without first serum iron. In each fold, numeric median imputation, categorical mode imputation, "
        "one-hot encoding, scaling, and unweighted logistic regression were fitted on the training "
        "partition only. Paired bootstrap resampling quantified confidence intervals for changes in "
        "out-of-fold AUROC and Brier score."
    ),
    82: (
        "The updated Figure 1 summarizes the overall workflow and the analysis sets displayed in the "
        "artwork. In MIMIC-IV, 1,568 admissions met the ≥48-hour ICU-stay screen, 1,487 also met the "
        "documented serum iron <60 µg/dL definition (215 endpoint events), and 1,344 formed the locked "
        "complete-feature cohort (200 events). The remaining 143 eligible admissions had 15 events."
    ),
    83: (
        "Table 1 describes all 1,487 eligible MIMIC-IV admissions by 28-day in-hospital outcome from "
        "hospital admission. Event prevalence was 14.5% (215/1,487). Compared with the complete-feature "
        "group, excluded patients were older and had higher first lactate but much shorter ICU stays and "
        "substantially less documented ventilation; their endpoint mortality was lower (10.5% vs 14.9%; "
        "Supplementary Table S11)."
    ),
    84: (
        "The corrected MIMIC-III cohort contained 799 admissions and 140 endpoint events after applying "
        "the stated ≥48-hour ICU-length-of-stay criterion. NIV documentation was missing for 58.6% of "
        "this cohort. F3 and NH contained 101 and 98 admissions, respectively."
    ),
    85: (
        "The reduced model yielded AUROC 0.810 in F3 and 0.898 in NH. The higher NH estimate should not "
        "be interpreted as superior transportability because NH was small and had high event prevalence "
        "(37/98); its calibration slope was 2.387, consistent with substantial probability miscalibration."
    ),
    87: "Figure 1. Study design, cohort flow, and reproducibility architecture",
    88: (
        "(A) Overview of data collection, cohort assembly, feature extraction, machine-learning "
        "development, internal and non-training evaluation, and the research web application."
    ),
    89: (
        "(B) Cohort allocation shown in the author-supplied schematic: MIMIC-IV training (n=1,075) and "
        "internal test (n=269) sets; corrected MIMIC-III temporal evaluation (n=799); and exploratory "
        "F3 (n=101) and NH (n=98) geographic evaluations of the reduced model. The sequential MIMIC-IV "
        "screening denominators (1,568 after ICU LOS ≥48 hours, 1,487 meeting the fixed iron criterion, "
        "and 1,344 in the complete locked matrix) are reported in the Methods, Results, and "
        "Supplementary Table S7. Figure created by the authors with BioRender.com."
    ),
    103: (
        "Table 1. Characteristics of all eligible MIMIC-IV admissions with documented biochemical "
        "hypoferremia, stratified by 28-day in-hospital outcome from hospital admission"
    ),
    104: (
        "Note: Table 1 uses the eligibility denominator (n=1,487; 215 events), not the complete locked "
        "modeling denominator. Continuous variables are shown as median (Q1, Q3) and categorical "
        "variables as n (%). Variables labelled first use the first recorded ICU-stay value; minimum, "
        "maximum, and mean values summarize the full ICU stay and are not admission-only measurements."
    ),
    130: "2. Serum iron within the documented-hypoferremia cohort",
    131: (
        "Within the 1,487 eligible MIMIC-IV admissions, the adjusted odds ratio for 28-day in-hospital "
        "mortality was 1.05 per 10 µg/dL higher first serum iron (95% CI 0.95–1.17; P=0.348) after "
        "adjustment for age, sex, race, SOFA score, Charlson comorbidity index, first lactate, invasive "
        "ventilation, and vasoactive therapy. The complete-case sensitivity analysis (n=879; 160 events) "
        "was concordant (OR 1.06; 95% CI 0.93–1.20; P=0.406). These estimates do not support an "
        "independent linear association in the selected cohort."
    ),
    132: (
        "A four-knot restricted cubic spline gave no evidence of nonlinearity (likelihood-ratio P=0.555; "
        "Figure 2). Adding first serum iron to the clinical covariate model changed five-fold out-of-fold "
        "AUROC by −0.0014 (95% CI −0.0075 to 0.0044) and Brier score by 0.0004 (95% CI "
        "−0.00005 to 0.00077). These findings support treating documented hypoferremia as a cohort-entry "
        "condition rather than the dominant model signal."
    ),
    134: "Figure 2. Adjusted association of first serum iron with mortality within documented hypoferremia",
    135: (
        "Adjusted odds-ratio curve and 95% confidence band for first serum iron among the 1,487 eligible "
        "MIMIC-IV admissions; the histogram shows the observed distribution. The curve is referenced "
        "to the cohort median. Covariates were age, sex, race, SOFA score, Charlson comorbidity index, "
        "first lactate, invasive ventilation, and vasoactive therapy. Restricted cubic spline knots "
        "were placed at 9, 20, 33, and 55 µg/dL (5th, 35th, 65th, and 95th percentiles); the "
        "nonlinearity likelihood-ratio test was P=0.555. No normal-iron comparator is shown because "
        "a reproducible comparator dataset and extraction code were not retained."
    ),
    153: "4. Model development, temporal evaluation, and comparative performance",
    154: (
        "In the MIMIC-IV internal test set (n=269; 40 events), the locked 18-predictor soft-voting "
        "benchmark achieved AUROC 0.840 (95% CI 0.767–0.913). Logistic regression achieved 0.846 and "
        "SVM 0.848; therefore, the ensemble did not demonstrate superiority over simpler or individual "
        "learners. It remains the benchmark only because it was prespecified."
    ),
    156: (
        "After reapplying ICU LOS ≥48 hours, the corrected MIMIC-III temporal cohort contained 799 "
        "admissions (140 events). AUROC was 0.705 for the 18-predictor benchmark and 0.715 for the "
        "8-feature derivative, compared with 0.840 and 0.796 internally. This temporal decline argues "
        "against assuming broad transportability."
    ),
    157: (
        "All current MIMIC-III tables, figures, performance estimates, and conclusions use the corrected "
        "n=799 temporal evaluation cohort."
    ),
    160: "Figure 4. Discrimination and calibration summary across evaluation cohorts",
    161: (
        "AUROC of the locked 18-predictor benchmark in the MIMIC-IV internal test set and corrected "
        "MIMIC-III temporal cohort. (B) AUROC of the post hoc 8-feature derivative in MIMIC-IV, corrected "
        "MIMIC-III, F3, and NH."
    ),
    162: (
        "Values are shown with cohort role and event counts. F3 and NH estimates are exploratory; their "
        "small sample sizes, different event prevalences, and calibration error preclude claims of "
        "superior geographic transportability."
    ),
    177: (
        "Table 3. Corrected evaluation of the locked 18-predictor and post hoc 8-feature soft-voting "
        "models in MIMIC-III after applying ICU LOS ≥48 hours (n=799; 140 events)"
    ),
    185: "5. Calibration and retrospective decision-curve analysis",
    186: (
        "Calibration was quantified using Brier score, calibration-in-the-large (CITL), and calibration "
        "slope. For the 18-predictor benchmark, internal values were 0.109, −0.898, and 1.705; in corrected "
        "MIMIC-III they were 0.141, −0.597, and 0.929. For the 8-feature derivative, the corresponding "
        "values were 0.118, −0.975, and 1.338 internally and 0.134, −0.323, and 0.855 in corrected "
        "MIMIC-III. The negative CITL values indicate systematic overprediction on average."
    ),
    187: (
        "Decision-curve analysis was retrospective and exploratory. Threshold probability represents "
        "the risk level at which the harm of a false-positive intervention is judged equivalent to the "
        "benefit of identifying a true event. We did not prespecify a clinical action or a deployment "
        "threshold; therefore, net-benefit curves should not be read as evidence of clinical utility."
    ),
    188: (
        "Small event counts, especially 40 internal-test events and the two Chinese cohorts, limit the "
        "precision of calibration estimates. The models require independent validation and local "
        "recalibration before individual predicted probabilities could be used."
    ),
    190: "Figure 5. Internal calibration and exploratory retrospective decision-curve analyses",
    191: (
        "(A) Calibration curves in the MIMIC-IV internal test set. (B) Decision curves across candidate "
        "thresholds. (C) Clinical-impact curve for the locked soft-voting benchmark. These graphical "
        "analyses complement, but do not replace, the intercept/slope/Brier metrics in Supplementary "
        "Table S10. No threshold was selected for clinical use."
    ),
    194: (
        "SHAP was used to describe how predictor values contributed to model output in the analyzed "
        "data. The largest contributions arose from general illness-severity and treatment-documentation "
        "variables, including oxygen saturation, ventilation documentation, platelet count, renal markers, "
        "and blood gases; serum iron was not a dominant feature."
    ),
    195: (
        "The NIV indicator captured documented use and may also reflect recording practices. Its high "
        "missingness in corrected MIMIC-III and importance in the original model motivated a sensitivity "
        "analysis excluding it. Removing NIV changed internal AUROC of the benchmark from 0.840 to 0.816 "
        "and gave AUROC 0.727 in corrected MIMIC-III."
    ),
    196: (
        "Local SHAP examples illustrate statistical contributions for individual records. They should "
        "not be interpreted as mechanisms, treatment effects, or explanations of why a patient died. "
        "References to a real t0 were removed because the retained export did not contain validated iron timestamps."
    ),
    198: (
        "(A) Global SHAP summary. (B) Mean absolute SHAP values. (C–D) Two illustrative local explanation "
        "plots. Color and direction show contributions to model output, not causal effects. The NIV field "
        "represents documentation and may be dataset-specific."
    ),
    205: "7. Performance and interpretation of the post hoc reduced model",
    206: (
        "The 8-feature model was developed post hoc from the broader development-data feature and SHAP "
        "results to support future implementation research. It was not prespecified, does not replace "
        "the benchmark, and its assessment on the internal test set may carry optimism. None of its "
        "features directly measures iron metabolism; its predictions primarily reflect general respiratory, "
        "renal, hematologic, and treatment-documentation severity."
    ),
    207: (
        "The reduced model achieved AUROC 0.796 internally, 0.715 in corrected MIMIC-III, 0.810 in F3, "
        "and 0.898 in NH. The geographically external estimates are exploratory because F3 and NH are "
        "small. Higher NH discrimination coexisted with substantial miscalibration (CITL 0.386; slope 2.387)."
    ),
    211: "Figure 7. Post hoc 8-feature model: explainability and cohort-level evaluation",
    212: (
        "(A) SHAP summary for the 8-feature derivative. (B) Cohort-level AUROC for MIMIC-IV, corrected "
        "MIMIC-III, F3, and NH. (C–D) Illustrative local SHAP explanations. The plots describe prediction "
        "within these datasets and do not establish causal importance or transportability."
    ),
    213: (
        "Across the four available calibration cohorts, 8-feature CITL/slope values were −0.975/1.338 "
        "(MIMIC-IV), −0.323/0.855 (corrected MIMIC-III), −0.336/1.351 (F3), and 0.386/2.387 (NH). "
        "These results replace earlier binned-regression slope estimates and show that recalibration is necessary."
    ),
    215: (
        "In the NIV-exclusion sensitivity analysis, the 7-feature derivative achieved AUROC 0.783 "
        "internally and 0.751 in corrected MIMIC-III. This result suggests that documentation artifacts "
        "may influence transportability and that the NIV variable should not be interpreted physiologically."
    ),
    216: (
        "Any future decision-curve or impact analysis must be tied to a prespecified clinical action, "
        "validated threshold, and prospective workflow. The present curves are exploratory."
    ),
    217: "8. Research-only software demonstration",
    218: (
        "The 8-feature interface is a research demonstration of model input/output and explanation "
        "rendering. It is not a medical device, has not undergone prospective impact evaluation, and is "
        "not approved or intended for bedside decision-making."
    ),
    220: (
        "Before any deployment, the model would require independent multicenter validation, local "
        "recalibration, clinical-governance review, usability testing, and prospective impact assessment."
    ),
    222: "Figure 8. Research-only web demonstration of the post hoc 8-feature model",
    223: (
        "Screenshot of the research interface. All inputs use model-native units (creatinine and BUN "
        "in mg/dL; non-invasive ventilation coded as documented use versus no documented use). This "
        "archived screenshot predates correction of the output label and therefore displays the complementary "
        "28-day survival probability; the current repository interface consistently reports 28-day mortality "
        "probability. SHAP summary, waterfall, and force plots are provided for research inspection. The "
        "interface is not a medical device, its probability has not been locally recalibrated for bedside use, "
        "and the display must not guide diagnosis, triage, or treatment."
    ),
    227: (
        "Documented biochemical hypoferremia is common during critical illness and may reflect cytokine- "
        "and hepcidin-mediated iron sequestration, reduced circulating transferrin iron, organ dysfunction, "
        "or treatment effects. A low serum iron value alone cannot distinguish inflammation-associated "
        "restriction from absolute iron-store depletion and should not be interpreted as a causal driver of mortality."
    ),
    228: (
        "Our additional analysis clarifies the scientific role of serum iron. Within patients already "
        "restricted to iron <60 µg/dL, first serum iron was neither independently associated with the "
        "endpoint nor incrementally discriminative. The model is therefore best understood as general "
        "severity risk stratification in an ICU population selected by documented biochemical hypoferremia, "
        "not an iron-driven mortality model."
    ),
    229: (
        "The 18-predictor benchmark discriminated internally but lost performance in the corrected "
        "same-institution temporal cohort. Similar temporal degradation is expected when measurement, "
        "documentation, case mix, and practice patterns change. The small F3 and NH studies provide "
        "exploratory geographic evidence for the reduced model only and do not establish universal transportability."
    ),
    230: (
        "SHAP concentrated prediction in oxygenation, ventilation documentation, platelets, blood gases, "
        "and renal markers. These features are clinically coherent indicators of severity, but SHAP "
        "importance is conditional on the fitted model and correlated inputs; it is not biological causality. "
        "The lack of iron variables in the reduced panel reinforces this distinction."
    ),
    231: (
        "The soft-voting ensemble was retained for protocol continuity, not superiority. Logistic regression "
        "matched or exceeded it internally, consistent with evidence that more complex machine learning "
        "does not automatically outperform regression in clinical prediction. Simpler models may be preferable "
        "after prospective comparison of calibration, maintenance burden, and transportability."
    ),
    232: (
        "Calibration results materially limit immediate application. Negative CITL in MIMIC-IV and corrected "
        "MIMIC-III indicated average overprediction, and slopes departed from one in several cohorts. Discrimination "
        "alone is insufficient for patient-level decisions; local recalibration and clinical-impact testing are required."
    ),
    233: (
        "Multiple imputation retained all eligible MIMIC-IV admissions and produced lower out-of-fold AUROC "
        "than the locked complete-feature internal result. The excluded admissions were not simply the sickest: "
        "they had lower mortality and shorter stays but different measurement and ventilation patterns. This "
        "still represents selection by data availability and limits generalizability."
    ),
    234: (
        "Other limitations include 160 training events for 18 locked predictors, leaving a modest effective "
        "sample size relative to model flexibility and a plausible risk of overfitting; post hoc derivation of "
        "the reduced panel; dataset-specific imputation and NIV documentation; full-ICU-stay summaries that "
        "prevent admission-time use and can incorporate information occurring late in the stay; small exploratory "
        "Chinese cohorts; and an in-hospital endpoint that classifies discharge before day 28 as survival without "
        "capturing post-discharge deaths."
    ),
    235: (
        "The retained export lacked authentic iron timestamps and a reproducible normal-iron comparator cohort. "
        "Accordingly, we removed t0 and normal-versus-low-iron claims rather than impute unavailable timing or "
        "reconstruct an unsupported comparator. Future work should begin from timestamped raw extraction code, "
        "prespecify a prediction time, validate in independent multicenter populations, and assess recalibration "
        "and impact prospectively."
    ),
    237: (
        "Among ICU patients selected by documented biochemical hypoferremia, interpretable models provided "
        "retrospective mortality-risk stratification but mainly captured general severity rather than an "
        "independent serum-iron signal. Corrected temporal evaluation, missing-data sensitivity, and calibration "
        "results support cautious research use only; prospective multicenter validation and local recalibration "
        "are required before clinical application."
    ),
}


def _save_with_replaced_media(doc: Document, out: Path, *, highlighted: bool) -> None:
    temp_docx = out.with_suffix(".tmp.docx")
    doc.save(temp_docx)
    fig2_tif = out.parent / "_figure2_r2.tiff"
    fig4_tif = out.parent / "_figure4_r2.tiff"
    with Image.open(REF / "figure2_within_hypoferremia_dose_response.png") as im:
        im.convert("RGB").save(fig2_tif, format="TIFF", compression="tiff_lzw", dpi=(300, 300))
    with Image.open(REF / "figure4_corrected_cohort_performance.png") as im:
        im.convert("RGB").save(fig4_tif, format="TIFF", compression="tiff_lzw", dpi=(300, 300))
    media = {
        "word/media/image1.png": UPDATED_FIGURE1,
        "word/media/image2.tiff": fig2_tif,
        "word/media/image4.tiff": fig4_tif,
        "word/media/image7.png": REF / "figure7_reduced_model_corrected.png",
    }
    with zipfile.ZipFile(temp_docx, "r") as src, zipfile.ZipFile(
        out, "w", compression=zipfile.ZIP_DEFLATED
    ) as dst:
        for item in src.infolist():
            data = media[item.filename].read_bytes() if item.filename in media else src.read(item.filename)
            dst.writestr(item, data)
    temp_docx.unlink()
    fig2_tif.unlink()
    fig4_tif.unlink()


def _set_inline_shape_alt_text(shape, title: str, description: str) -> None:
    """Add accessibility metadata without changing the visible artwork."""
    doc_pr = shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def build_manuscript(highlighted: bool) -> Path:
    src = R1 / "Hypoferremia_Manuscript_R1_Clean.docx"
    doc = Document(src)
    _set_document_metadata(
        doc,
        NEW_TITLE,
        "Round 2 highlighted manuscript" if highlighted else "Round 2 clean manuscript",
    )
    # Paragraphs 62–63 are blank insertion anchors in R1. Give them the same
    # geometry and hierarchy as the existing Methods heading/body before adding
    # the new serum-iron subsection.
    _copy_paragraph_properties(doc.paragraphs[46], doc.paragraphs[62])
    set_paragraph(
        doc.paragraphs[62],
        MANUSCRIPT_REPLACEMENTS[62],
        highlight=highlighted,
        bold=True,
        font_size=12.0,
    )
    _copy_paragraph_properties(doc.paragraphs[56], doc.paragraphs[63])
    set_paragraph(
        doc.paragraphs[63],
        MANUSCRIPT_REPLACEMENTS[63],
        highlight=highlighted,
        bold=False,
        font_size=12.0,
    )
    for idx, text in MANUSCRIPT_REPLACEMENTS.items():
        if idx in {62, 63}:
            continue
        set_paragraph(doc.paragraphs[idx], text, highlight=highlighted)
    # Figure 6 and its title share one paragraph in the R1 source. Revise only
    # the textual run so the inline artwork remains embedded.
    set_text_preserving_drawing(
        doc.paragraphs[197],
        "Figure 6. SHAP-based description of the locked soft-voting model",
        highlight=highlighted,
    )

    references = next(p for p in doc.paragraphs if p.text.strip() == "References")
    heading_template = next(
        p for p in doc.paragraphs if p.text.strip() == "Conflicts of Interest"
    )
    body_template = next(
        p
        for p in doc.paragraphs
        if p.text.strip().startswith("MIMIC-III and MIMIC-IV were accessed")
    )
    availability = references.insert_paragraph_before()
    _copy_paragraph_properties(body_template, availability)
    set_paragraph(
        availability,
        "The prepared-modeling-table analysis code, locked model and preprocessing objects, and aggregate "
        "Round 2 outputs are available at https://github.com/linzixuan45/hypoferremia-icu-mortality. "
        "Release v1.1.0 is prepared for Zenodo deposit. At the time this marked revision was prepared, the "
        "record had not yet been published; the final Zenodo DOI must be inserted before resubmission. "
        "Public files exclude patient-level MIMIC or local-cohort records, outcome labels, and "
        "individual prediction probabilities. The public repository begins from authorized prepared "
        "modeling tables and does not constitute a complete raw-database-to-modeling-table extraction pipeline.",
        highlight=highlighted,
        bold=False,
        font_size=12.0,
    )
    availability_heading = availability.insert_paragraph_before()
    _copy_paragraph_properties(heading_template, availability_heading)
    set_paragraph(
        availability_heading,
        "Code and data availability",
        highlight=highlighted,
        bold=True,
        font_size=12.0,
    )
    availability_heading.paragraph_format.keep_with_next = True

    replace_table(doc, 0, load_table1(), highlight=highlighted)

    corrected = pd.read_csv(REF / "r2_corrected_mimic3_evaluation.csv")
    corrected_display = pd.DataFrame(
        {
            "Model": ["18-predictor benchmark", "8-feature post hoc derivative"],
            "n": corrected["n"].astype(int),
            "Events": corrected["Events"].astype(int),
            "AUROC": corrected["AUROC"].map(lambda x: f"{x:.3f}"),
            "Brier": corrected["Brier"].map(lambda x: f"{x:.3f}"),
            "CITL": corrected["Calibration_in_the_large"].map(lambda x: f"{x:.3f}"),
            "Calibration slope": corrected["Calibration_slope"].map(lambda x: f"{x:.3f}"),
        }
    )
    replace_table(doc, 2, df_rows(corrected_display), highlight=highlighted)

    # Preserve the R1 A4 page geometry. Revised artwork is fit to the original
    # usable width rather than changing margins across the entire manuscript.
    for shape_i, shape in enumerate(doc.inline_shapes):
        if shape_i == 0:
            shape.width, shape.height = Inches(5.75), Inches(5.75)
        elif shape_i == 1:
            shape.width, shape.height = Inches(5.75), Inches(4.92)
        elif shape_i == 3:
            shape.width, shape.height = Inches(5.75), Inches(2.56)
        elif shape_i == 6:
            shape.width, shape.height = Inches(5.75), Inches(3.70)
        if shape_i in {0, 1, 3, 6}:
            # The R1 placeholders carried image-specific crop rectangles. Reusing
            # those crop settings would hide parts of the replacement figures.
            for node in list(shape._inline.iter()):
                if node.tag == qn("a:srcRect"):
                    node.getparent().remove(node)
        _set_inline_shape_alt_text(
            shape,
            f"Figure {shape_i + 1}",
            (
                "Study figure; detailed content and interpretation are provided in the "
                f"Figure {shape_i + 1} legend."
            ),
        )

    # Enforce one publication-style table system across unchanged and revised tables.
    for table in doc.tables:
        _apply_three_line_table(table)

    out = R2 / (
        "Hypoferremia_Manuscript_R2_Highlighted.docx"
        if highlighted
        else "Hypoferremia_Manuscript_R2_Clean.docx"
    )
    _save_with_replaced_media(doc, out, highlighted=highlighted)
    return out


def build_supplement(highlighted: bool) -> Path:
    doc = Document(R1 / "Hypoferremia_Supplementary_R1_Clean.docx")
    _set_document_metadata(
        doc,
        f"Supplementary Tables — {NEW_TITLE}",
        "Round 2 highlighted supplementary material"
        if highlighted
        else "Round 2 clean supplementary material",
    )
    changes = {
        0: "Supplementary Tables — Round 2",
        1: NEW_TITLE,
        4: "Supplementary Table S1. Corrected MIMIC-III temporal evaluation after applying ICU LOS ≥48 hours",
        5: (
            "The corrected analytic denominator is n=799 with 140 deaths within 28 days of hospital "
            "admission. Earlier n=1,046 results did not apply the stated LOS criterion and are superseded."
        ),
        7: "Supplementary Table S2. Characteristics of the Third Xiangya (F3) cohort.",
        10: "Supplementary Table S3. Characteristics of the Nanhua Hospital (NH) cohort.",
        13: "Supplementary Table S4. Hyperparameters and tuning strategies for all 12 models.",
        19: "Supplementary Table S6. Locked predictor missingness and actual observation windows in eligible MIMIC-IV admissions.",
        20: (
            "Suffix definitions: _1st is the first recorded ICU-stay value; _min, _max, and _mean summarize "
            "the full ICU stay. los_icu_to_t0 is a synthetic legacy proxy and not an authentic laboratory timestamp."
        ),
        21: "Supplementary Table S7. Corrected patient flow and endpoint counts.",
        22: (
            "Eligibility (n=1,487) is separated from complete-matrix inclusion (n=1,344). The difference "
            "is 143 admissions with 15 endpoint events."
        ),
        23: (
            "All current temporal-evaluation tables use the corrected ≥48-hour MIMIC-III cohort (n=799). "
            "Superseded temporal estimates are not included in the current sensitivity tables."
        ),
        24: (
            "Supplementary Table S8. Class-weight primary versus random-undersampling sensitivity in the "
            "MIMIC-IV internal test set (locked analysis)."
        ),
        25: "Supplementary Table S9. Five-imputation, five-fold out-of-fold sensitivity in all eligible MIMIC-IV admissions.",
        26: (
            "The multiple-imputation analysis retained n=1,487 and used only aggregate metrics for public "
            "reporting; patient-level outcomes and probabilities remain protected."
        ),
        27: "Supplementary Table S10. Standard calibration metrics across available validation cohorts.",
    }
    for idx, text in changes.items():
        set_paragraph(doc.paragraphs[idx], text, highlight=highlighted)

    def compact_metrics(frame: pd.DataFrame) -> pd.DataFrame:
        return pd.DataFrame(
            {
                "Analysis": frame["Analysis"],
                "Cohort": frame["Cohort"],
                "n (events)": [
                    f"{int(n)} ({int(e)})" for n, e in zip(frame["n"], frame["Events"])
                ],
                "AUROC": frame["AUROC"].map(lambda x: f"{x:.3f}"),
                "Brier (95% CI)": [
                    f"{v:.3f} ({lo:.3f}–{hi:.3f})"
                    for v, lo, hi in zip(frame["Brier"], frame["Brier_95CI_lo"], frame["Brier_95CI_hi"])
                ],
                "CITL (95% CI)": [
                    f"{v:.3f} ({lo:.3f}–{hi:.3f})"
                    for v, lo, hi in zip(
                        frame["Calibration_in_the_large"],
                        frame["CITL_95CI_lo"],
                        frame["CITL_95CI_hi"],
                    )
                ],
                "Slope (95% CI)": [
                    f"{v:.3f} ({lo:.3f}–{hi:.3f})"
                    for v, lo, hi in zip(
                        frame["Calibration_slope"],
                        frame["Slope_95CI_lo"],
                        frame["Slope_95CI_hi"],
                    )
                ],
            }
        )

    c_m3 = pd.read_csv(REF / "r2_corrected_mimic3_evaluation.csv")
    replace_table(doc, 0, df_rows(compact_metrics(c_m3)), highlight=highlighted)

    # Synchronize the duplicate checklist in the supplement with stable current
    # section titles; numeric subsection labels from R1 no longer exist.
    supplement_checklist = doc.tables[4]
    checklist_updates = {
        5: ("Methods — Data sources and study population", "Yes"),
        6: ("Methods — Study population; Figure 1; Supplementary Table S7", "Yes"),
        7: ("Methods — Variables, preprocessing, and outcome definition", "Yes"),
        8: ("Methods — Predictor observation window; Supplementary Tables S4 and S6", "Yes"),
        9: ("Methods — Study population; Results — Study population", "Partly reported"),
        10: (
            "Methods — Feature selection and data preprocessing; Supplementary Tables S6, S9, and S11",
            "Yes",
        ),
        11: ("Methods — Machine learning model development and evaluation", "Yes"),
        12: ("Methods — Machine learning model development; Supplementary Table S4", "Yes"),
        13: ("Methods — Feature selection and data preprocessing", "Yes"),
        14: (
            "Methods — Machine learning model development; Supplementary Tables S4 and S15",
            "Partly reported",
        ),
        15: (
            "Methods — Machine learning model development; repository requirements files",
            "Partly reported",
        ),
        18: ("Results — Calibration; Supplementary Table S10", "Yes"),
        23: ("Supplement; Code and Data Availability; public repository", "Yes"),
        25: (
            "Code and Data Availability; public repository; Zenodo DOI to be inserted before resubmission",
            "Yes",
        ),
    }
    for row_index, (location, status) in checklist_updates.items():
        set_paragraph(
            supplement_checklist.rows[row_index].cells[3].paragraphs[0],
            location,
            highlight=highlighted,
        )
        set_paragraph(
            supplement_checklist.rows[row_index].cells[4].paragraphs[0],
            status,
            highlight=highlighted,
        )

    windows = pd.read_csv(REF / "r2_locked_matrix_missingness_and_windows.csv")
    windows["Missing_pct"] = windows["Missing_pct"].map(lambda x: f"{x:.1f}%")
    windows = windows.rename(
        columns={
            "Missing_n": "Missing, n",
            "Missing_pct": "Missing, %",
            "Observation_window": "Observation window",
        }
    )
    replace_table(doc, 5, df_rows(windows), highlight=highlighted)
    flow = pd.read_csv(REF / "r2_patient_flow.csv")
    flow = flow.rename(
        columns={
            "Hospital_deaths": "Hospital deaths, n",
            "Hospital_mortality_pct": "Hospital mortality, %",
            "Deaths_within_28d_of_hospital_admission": (
                "28-day in-hospital deaths, n"
            ),
            "28d_in_hospital_mortality_pct": (
                "28-day in-hospital mortality, %"
            ),
        }
    )
    replace_table(doc, 6, df_rows(flow), highlight=highlighted)
    # The retained R1 table included obsolete MIMIC-III n=1,046 rows. Keep the
    # auditable internal class-imbalance sensitivity and remove those superseded
    # temporal rows instead of relabelling results from an ineligible cohort.
    s8_rows = [
        [cell.text for cell in row.cells]
        for row_index, row in enumerate(doc.tables[7].rows)
        if row_index == 0 or "MIMIC-IV test" in row.cells[1].text
    ]
    replace_table(doc, 7, s8_rows, highlight=highlighted)
    mi = pd.read_csv(REF / "r2_multiple_imputation_sensitivity.csv")
    replace_table(doc, 8, df_rows(compact_metrics(mi)), highlight=highlighted)
    calibration = pd.read_csv(REF / "r2_standard_calibration_metrics.csv")
    replace_table(doc, 9, df_rows(compact_metrics(calibration)), highlight=highlighted)

    # Reuse the R1 caption and note geometry for every newly added table.
    caption_template = doc.paragraphs[27]
    note_template = doc.paragraphs[23]
    included = pd.read_csv(REF / "r2_included_vs_excluded.csv")
    included = included.rename(
        columns={
            "Included_complete_case": "Included complete-case",
            "Excluded_for_missingness": "Excluded for missingness",
            "Absolute_SMD": "Absolute SMD",
            "P_value": "P value",
        }
    )
    append_table(
        doc,
        "Supplementary Table S11. Eligible MIMIC-IV admissions included in versus excluded from the complete locked matrix.",
        df_rows(included),
        "Absolute standardized mean differences are descriptive; P values are not corrected for multiplicity.",
        highlight=highlighted,
        caption_template=caption_template,
        note_template=note_template,
    )
    niv = pd.read_csv(REF / "r2_niv_exclusion_sensitivity.csv")
    append_table(
        doc,
        "Supplementary Table S12. Sensitivity analysis excluding the NIV documentation variable.",
        df_rows(compact_metrics(niv)),
        "The indicator may encode both treatment and documentation availability; it is not interpreted as a causal exposure.",
        highlight=highlighted,
        caption_template=caption_template,
        note_template=note_template,
    )
    iron = pd.read_csv(REF / "r2_iron_adjusted_effect.csv")
    iron_display = pd.DataFrame(
        {
            "Analysis": iron["Analysis"],
            "n (events)": [
                f"{int(n)} ({int(e)})" for n, e in zip(iron["n"], iron["Events"])
            ],
            "OR per 10 µg/dL (95% CI)": [
                "—" if pd.isna(value) else f"{value:.3f} ({lo:.3f}–{hi:.3f})"
                for value, lo, hi in zip(
                    iron["Adjusted_OR"], iron["OR_95CI_lo"], iron["OR_95CI_hi"]
                )
            ],
            "P value": iron["P_value"].map(lambda x: f"{x:.3f}"),
            "Test interpretation": [
                "Unadjusted linear association",
                "Adjusted linear association",
                "Adjusted complete-case sensitivity",
                "Restricted cubic spline nonlinearity",
            ],
            "RCS knots (µg/dL)": ["9, 20, 33, 55"] * len(iron),
        }
    )
    append_table(
        doc,
        "Supplementary Table S13. First serum iron: adjusted association and nonlinearity within documented hypoferremia.",
        df_rows(iron_display),
        (
            "The adjusted model included age, sex, race (unordered; White reference), SOFA score, "
            "Charlson comorbidity index, first lactate, invasive ventilation, and vasoactive therapy. "
            "The primary analysis used numeric median and categorical mode imputation; the complete-case "
            "sensitivity retained n=879. Restricted cubic spline knots were placed at 9, 20, 33, and "
            "55 µg/dL (5th, 35th, 65th, and 95th percentiles)."
        ),
        highlight=highlighted,
        caption_template=caption_template,
        note_template=note_template,
    )
    incremental = pd.read_csv(REF / "r2_iron_incremental_value.csv")
    incremental_display = pd.DataFrame(
        {
            "Model": incremental["Model"],
            "n (events)": [
                f"{int(n)} ({int(e)})" for n, e in zip(incremental["n"], incremental["Events"])
            ],
            "OOF AUROC": incremental["OOF_AUROC"].map(lambda x: f"{x:.3f}"),
            "OOF Brier": incremental["OOF_Brier"].map(lambda x: f"{x:.3f}"),
            "ΔAUROC (95% CI)": [
                "Reference"
                if model == "Clinical covariates only"
                else f"{v:.4f} ({lo:.4f}–{hi:.4f})"
                for model, v, lo, hi in zip(
                    incremental["Model"],
                    incremental["Delta_AUROC_vs_clinical"],
                    incremental["Delta_AUROC_95CI_lo"],
                    incremental["Delta_AUROC_95CI_hi"],
                )
            ],
            "ΔBrier (95% CI)": [
                "Reference"
                if model == "Clinical covariates only"
                else f"{v:.4f} ({lo:.4f}–{hi:.4f})"
                for model, v, lo, hi in zip(
                    incremental["Model"],
                    incremental["Delta_Brier_vs_clinical"],
                    incremental["Delta_Brier_95CI_lo"],
                    incremental["Delta_Brier_95CI_hi"],
                )
            ],
        }
    )
    append_table(
        doc,
        "Supplementary Table S14. Incremental predictive value of first serum iron.",
        df_rows(incremental_display),
        (
            "Values are paired five-fold out-of-fold estimates in all 1,487 eligible MIMIC-IV "
            "admissions. Numeric and categorical imputation, one-hot encoding, scaling, and unweighted "
            "logistic regression were fitted within each training fold."
        ),
        highlight=highlighted,
        caption_template=caption_template,
        note_template=note_template,
    )
    pathway = [
        ["Model", "Status", "Predictors", "Development role", "Evaluation role", "Interpretation"],
        [
            "Locked benchmark",
            "Prespecified benchmark",
            "18",
            "MIMIC-IV complete matrix",
            "MIMIC-IV internal; corrected MIMIC-III temporal",
            "Primary protocol-continuity model; not superior to logistic regression",
        ],
        [
            "Reduced derivative",
            "Post hoc",
            "8",
            "Development-data feature/SHAP review",
            "MIMIC-IV internal; corrected MIMIC-III; exploratory F3 and NH",
            "Implementation-oriented research derivative; performance may be optimistic",
        ],
    ]
    append_table(
        doc,
        "Supplementary Table S15. Development pathway and intended role of the two models.",
        pathway,
        highlight=highlighted,
        caption_template=caption_template,
    )
    # Apply the same three-line convention to every original and newly added table.
    for table in doc.tables:
        _apply_three_line_table(table)
    out = R2 / (
        "Hypoferremia_Supplementary_R2_Highlighted.docx"
        if highlighted
        else "Hypoferremia_Supplementary_R2_Clean.docx"
    )
    doc.save(out)
    return out


def build_tripod() -> Path:
    doc = Document(R1 / "Hypoferremia_TRIPOD_AI_Checklist_R1.docx")
    _set_document_metadata(doc, f"TRIPOD+AI Checklist — {NEW_TITLE}", "Round 2 reporting checklist")
    set_paragraph(doc.paragraphs[0], "Completed TRIPOD+AI Checklist — Round 2 Updated")
    set_paragraph(
        doc.paragraphs[1],
        (
            "Round 2 audit: the checklist was re-reviewed after correcting the predictor observation "
            "window, outcome time origin, cohort denominators, missing-data sensitivity, temporal "
            "evaluation, calibration reporting, reduced-model status, and computational reproducibility."
        ),
    )
    table = doc.tables[0]
    location_updates = {
        5: "Methods — Data sources and study population",
        6: "Methods — Study population; Figure 1; Supplementary Table S7",
        7: "Methods — Variables, preprocessing, and outcome definition",
        8: "Methods — Predictor observation window; Supplementary Tables S4 and S6",
        9: "Methods — Study population; Results — Study population",
        10: "Methods — Feature selection and data preprocessing; Supplementary Tables S6, S9, and S11",
        11: "Methods — Machine learning model development and evaluation",
        12: "Methods — Machine learning model development; Supplementary Table S4",
        13: "Methods — Feature selection and data preprocessing",
        14: "Methods — Machine learning model development; Supplementary Tables S4 and S15",
        15: "Methods — Machine learning model development; repository requirements files",
    }
    for row_index, location in location_updates.items():
        table.rows[row_index].cells[4].text = location
    table.rows[18].cells[4].text = "Results — Calibration; Supplementary Table S10"
    table.rows[18].cells[5].text = (
        "Reported. Brier score, calibration-in-the-large, and calibration slope with confidence "
        "intervals are reported for available evaluation cohorts."
    )
    table.rows[23].cells[4].text = "Supplement; Code and Data Availability; public repository"
    table.rows[23].cells[5].text = (
        "Reported. Prepared-table analysis code, locked model/preprocessing objects, aggregate Round 2 "
        "outputs, and versioned archive metadata are supplied. Protected patient-level data and a complete "
        "raw-database extraction pipeline are not redistributed."
    )
    table.rows[25].cells[4].text = (
        "Code and Data Availability; https://github.com/linzixuan45/hypoferremia-icu-mortality"
    )
    table.rows[25].cells[5].text = (
        "Reported. Public code and model objects are available; release v1.1.0 is prepared for Zenodo, "
        "and the final DOI must be inserted before resubmission after the corresponding author publishes "
        "the archive."
    )
    for checklist_table in doc.tables:
        _apply_three_line_table(checklist_table, font_size=7.5)
    out = R2 / "Hypoferremia_TRIPOD_AI_Checklist_R2.docx"
    doc.save(out)
    return out


def _blank_document_from_template(path: Path) -> Document:
    """Preserve the R1 section, styles, fonts, and document settings; remove its content."""
    doc = Document(path)
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    return doc


def _source_paragraphs(name: str) -> list[str]:
    return [p.text for p in Document(REVIEW_SOURCES / name).paragraphs]


def _join_inclusive(paragraphs: list[str], start: int, end: int) -> str:
    return "\n\n".join(paragraphs[start : end + 1])


def _add_response_heading(doc: Document, text: str, *, reviewer: bool = False) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    run.bold = True
    _set_times_new_roman(run, size_pt=16 if reviewer else 12)


def _add_response(doc: Document, comment: str, response: str, changes: str) -> None:
    """Match R1: verbatim black italic comment, then blue response and location."""
    for source_paragraph in comment.split("\n\n"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(source_paragraph)
        run.italic = True
        run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        _set_times_new_roman(run, size_pt=11)

    if not response.lstrip().startswith(("We thank", "We sincerely thank")):
        raise ValueError(f"Every response must open with thanks: {response[:80]}")
    for response_i, response_paragraph in enumerate(response.split("\n\n")):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if response_i == 0:
            label = p.add_run("Response: ")
            label.bold = True
        p.add_run(response_paragraph)
        for run in p.runs:
            run.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)
            _set_times_new_roman(run, size_pt=11)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    label = p.add_run("Manuscript/Supplementary revision (highlighted in yellow): ")
    label.bold = True
    p.add_run(changes)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)
        _set_times_new_roman(run, size_pt=11)


def build_response_letter() -> Path:
    doc = _blank_document_from_template(R1 / "Response_to_Reviewers_R1.docx")
    _set_document_metadata(doc, f"Response to Reviewers — {NEW_TITLE}", "Round 2 point-by-point response")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Response to Reviewers")
    r.bold = True
    _set_times_new_roman(r, size_pt=16)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _set_times_new_roman(
        p.add_run("Submission ID: 9974989f-eff3-4bd0-a4dc-a3d1280eb44d"),
        size_pt=11,
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _set_times_new_roman(p.add_run(f"Manuscript Title: {NEW_TITLE}"), size_pt=11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    introduction = p.add_run(
        "We thank the Editor and Reviewers for their careful assessment. Before revising the manuscript, "
        "we completed a code-and-data audit and the requested supplementary statistical analyses. The "
        "revision corrects cohort flow and temporal evaluation, reports calibration and missing-data "
        "sensitivities, quantifies the adjusted and incremental contribution of serum iron, replaces the "
        "unreproducible comparator analysis, and narrows the scientific claim to retrospective mortality-risk "
        "stratification among ICU patients selected by documented biochemical hypoferremia."
    )
    _set_times_new_roman(introduction, size_pt=11)

    editor = _source_paragraphs("editor_and_portal.docx")[1]
    code_comment = editor[editor.index("Please note that") : editor.index("\n\nEditor Comments")]
    editor_comment = editor[
        editor.index('"The manuscript is generally well written')
        : editor.index("\n-Yusuf Secgin")
    ]
    inhouse_comment = editor[
        editor.index("Please could you clarify where the images")
        : editor.index("\n\nWe recommend submitting")
    ]

    _add_response_heading(doc, "Editor and In-House Editorial Comments", reviewer=True)
    _add_response_heading(doc, "Code-deposition requirement")
    _add_response(
        doc,
        code_comment,
        (
            "We thank the Editor for emphasizing this essential reproducibility requirement. We conducted "
            "a file-level audit of the analysis materials and prepared release v1.1.0 containing the public "
            "prepared-table analysis pipeline, the retained iron/timing SQL template, locked preprocessing "
            "and model objects, aggregate Round 2 outputs, CITATION.cff, .zenodo.json, a model card, and "
            "versioned release notes. These materials reproduce the reported analyses from authorized "
            "prepared modeling tables. We state explicitly that the repository does not contain a complete "
            "raw-clinical-database-to-modeling-table extraction pipeline, because that material was not "
            "retained and must not be implied by the availability statement.\n\n"
            "The Zenodo record has not yet been published. Consequently, the manuscript and response package "
            "do not state, cite, or imply an existing DOI. The corresponding author will verify and publish "
            "the staged archive and insert the assigned DOI before final resubmission. Patient-level MIMIC and local-cohort "
            "records, observed outcome labels, and individual prediction probabilities are excluded from the "
            "public release because redistribution would conflict with the governing data-use agreements. "
            "This distinction between public computational artifacts and controlled clinical data is now "
            "stated explicitly."
        ),
        (
            "The Code and Data Availability statements were separated and expanded. The public release "
            "metadata, model card, Zenodo deposit instructions, and release checklist were updated to identify "
            "what is supplied, what remains controlled, and that the DOI must be inserted after the staged "
            "archive is published and before final resubmission."
        ),
    )
    _add_response_heading(doc, "Editor Comment")
    _add_response(
        doc,
        editor_comment,
        (
            "We thank the Editor for the constructive assessment and for identifying the need to better "
            "connect statistical performance with the requirements for clinical translation. We revised the "
            "manuscript to distinguish retrospective discrimination from clinical usefulness and now present "
            "the internal AUROC of 0.840 together with the decline to 0.705 in the corrected MIMIC-III temporal "
            "cohort. We also report Brier score, calibration-in-the-large, and calibration slope, which show "
            "that acceptable ranking performance does not imply calibrated absolute-risk estimates.\n\n"
            "We further discuss differences in case mix, laboratory measurement, documentation, and treatment "
            "practice as plausible contributors to transportability loss. Throughout the manuscript, the model "
            "is framed as a retrospective research instrument; prospective multicenter validation, local "
            "recalibration, governance review, and formal impact assessment are specified as prerequisites "
            "before clinical use. We audited the title, objectives, results, discussion, conclusions, figure "
            "legends, and web-demonstration language to remove any implication of established clinical utility."
        ),
        (
            "The Title, Abstract, Methods, Results, Discussion, Conclusions, Figure 4 and Figure 7 legends, "
            "and research-demonstration description were revised to align performance reporting with these "
            "clinical-translation limitations."
        ),
    )
    _add_response_heading(doc, "In-House Editorial Comment")
    _add_response(
        doc,
        inhouse_comment,
        (
            "We thank the In-House Editorial Team for requesting explicit clarification of the visual sources "
            "and publication rights. Figure 1 is the updated author-prepared figure supplied for this revision "
            "and was created by the authors with BioRender.com. The figure legend now states this provenance, "
            "and the submission package includes a separate source-and-rights declaration. No third-party "
            "photograph or externally reproduced scientific illustration is presented as original material.\n\n"
            "For editorial transparency, the declaration instructs the corresponding author to confirm "
            "authorship with the journal and to retain the applicable BioRender publication-license or "
            "permission record. The newest artwork itself displays the corrected MIMIC-III temporal "
            "evaluation cohort (n=799), matching the Methods, Results, Figure 4, Table 3, and Supplementary "
            "Table S1. Its legend also identifies the MIMIC-IV training and internal-test sets and states "
            "where the sequential MIMIC-IV screening denominators are reported."
        ),
        (
            "Figure 1 was replaced with the updated BioRender export. Its legend and "
            "FIGURE1_SOURCE_DECLARATION.md now document authorship, BioRender provenance, rights handling, "
            "and the corrected cohort roles and denominators."
        ),
    )

    r3p = _source_paragraphs("reviewer3.docx")
    r3_items = [
        (
            "Reviewer 3 - Overall Assessment",
            _join_inclusive(r3p, 3, 4),
            (
                "We thank the Reviewer for the careful reassessment and for recognizing both the revisions "
                "already made and the remaining methodological concerns. We treated these concerns as requiring "
                "a complete reproducibility review rather than isolated textual changes. The audit identified "
                "three issues with direct interpretive consequences: conflation of eligibility and complete-"
                "matrix denominators, use of timing language unsupported by authentic laboratory timestamps, "
                "and a normal-iron comparison that could not be reconstructed from the retained data and code.\n\n"
                "We corrected the cohort flow, redefined the model as retrospective whole-stay risk "
                "stratification, removed the unreproducible comparator analysis, and performed the requested "
                "missing-data, calibration, transportability, NIV-exclusion, and serum-iron analyses. The "
                "revised claims are therefore narrower and correspond only to results that can be traced to "
                "the audited analysis files."
            ),
            (
                "The manuscript, Supplementary Information, Figures 1–2 and 4–7, TRIPOD+AI checklist, code "
                "documentation, and release materials were cross-audited and revised as specified in the "
                "point-by-point responses below."
            ),
        ),
        (
            "Reviewer 3 - Major Concern 1",
            r3p[6],
            (
                "We thank the Reviewer for detecting the apparent discrepancy in participant flow and for "
                "requesting a quantitative assessment of potential complete-case bias. The earlier diagram "
                "did not adequately distinguish the sequential denominators. The audited flow is: 1,931 "
                "admissions before the length-of-stay filter; 1,568 after ICU length of stay ≥48 hours; "
                "1,487 meeting serum iron <60 µg/dL; and 1,344 with a complete locked 18-predictor matrix. "
                "Accordingly, 143 eligible admissions—not 224—were excluded for matrix incompleteness, and "
                "15—not 87—experienced the endpoint. Mortality was 10.5% among excluded admissions and 14.9% "
                "among included admissions.\n\n"
                "To assess whether complete-matrix restriction materially altered performance, we performed "
                "five-imputation, five-fold out-of-fold analysis in all 1,487 eligible MIMIC-IV admissions. "
                "The resulting AUROC was 0.808, Brier score 0.113, calibration-in-the-large −0.918, and "
                "calibration slope 1.453. These results support broadly similar ranking performance while "
                "also documenting residual probability miscalibration; they do not eliminate the possibility "
                "of selection bias from variables unavailable for imputation or from the original testing "
                "process."
            ),
            (
                "Figure 1 and its legend, Table 1, the cohort-selection and missing-data Methods, the Results "
                "and limitations, and Supplementary Tables S7, S9, and S11 now report the sequential "
                "denominators, event counts, included-versus-excluded comparison, and imputation sensitivity."
            ),
        ),
        (
            "Reviewer 3 - Major Concern 2",
            r3p[7],
            (
                "We thank the Reviewer for emphasizing that temporal degradation is central to the study's "
                "interpretation. Reapplying the stated ICU length-of-stay criterion corrected the MIMIC-III "
                "analytic cohort from the earlier export count to n=799 with 140 events. The locked "
                "18-predictor benchmark achieved AUROC 0.705 in this corrected cohort, compared with 0.840 "
                "in the MIMIC-IV internal test set. This decline is now presented as evidence of limited "
                "temporal transportability rather than as a successful external replication.\n\n"
                "Because MIMIC-III and MIMIC-IV derive from the same institution in different eras, we now "
                "describe this analysis as same-institution temporal evaluation. The two small Chinese "
                "cohorts are separately labelled exploratory geographic evaluations of the post hoc reduced "
                "model. This terminology prevents temporal and geographic validation from being conflated."
            ),
            (
                "The Title, cohort definitions in Methods, temporal-evaluation Results, Figure 4 and its "
                "legend, Table 3, Discussion, and Supplementary Table S1 were revised to use the corrected "
                "n=799 cohort and the same-institution temporal-evaluation terminology."
            ),
        ),
        (
            "Reviewer 3 - Major Concern 3",
            r3p[8],
            (
                "We thank the Reviewer for identifying calibration as necessary for evaluating probability "
                "predictions. We added Brier score, calibration-in-the-large (CITL), and calibration slope "
                "with confidence intervals for each cohort in which individual predictions were available. "
                "For the benchmark model, the internal values were Brier 0.109, CITL −0.898, and slope "
                "1.705; in corrected MIMIC-III they were 0.141, −0.597, and 0.929, respectively.\n\n"
                "These results indicate systematic overprediction on the log-odds scale in both cohorts and "
                "a change in calibration slope across time. We now interpret AUROC together with calibration "
                "rather than as a stand-alone measure of deployability, and explicitly note that estimates "
                "from the internal test set are imprecise because only 40 endpoint events were observed."
            ),
            (
                "The statistical-analysis Methods now define all calibration measures; the Results and "
                "Table 3 report the benchmark estimates; Supplementary Table S10 provides cohort-specific "
                "estimates and confidence intervals; and the Discussion explains the need for recalibration."
            ),
        ),
        (
            "Reviewer 3 - Major Concern 4",
            r3p[9],
            (
                "We thank the Reviewer for challenging the implied advantage of the soft-voting ensemble. "
                "We agree that the available results do not demonstrate superiority: logistic regression "
                "(AUROC 0.846) and the support-vector machine (AUROC 0.848) were numerically similar to or "
                "higher than the ensemble (AUROC 0.840) in the internal test set. We therefore retain the "
                "ensemble only as the locked protocol-continuity benchmark, not as the preferred clinical "
                "model.\n\n"
                "The revised Discussion states that model selection should consider calibration stability, "
                "transportability, computational maintenance, and interpretability in addition to small "
                "differences in internal AUROC. A simpler model may ultimately be preferable, but that choice "
                "requires prospectively specified head-to-head evaluation rather than post hoc ranking."
            ),
            (
                "The model-development Methods, algorithm-comparison Results, and Discussion were revised to "
                "identify the ensemble as a prespecified benchmark and to remove claims of demonstrated "
                "performance superiority."
            ),
        ),
        (
            "Reviewer 3 - Major Concern 5",
            r3p[10],
            (
                "We thank the Reviewer for identifying non-invasive ventilation (NIV) as a potentially "
                "documentation-sensitive predictor. In the corrected MIMIC-III cohort, NIV was unavailable "
                "for 468 of 799 admissions (58.6%), raising the possibility that the fitted model learned "
                "dataset-specific documentation or treatment patterns rather than a portable physiologic "
                "signal.\n\n"
                "We therefore repeated evaluation after excluding NIV. The 17-feature benchmark achieved "
                "AUROC 0.816 internally and 0.727 in corrected MIMIC-III; the corresponding 7-feature "
                "derivative achieved 0.783 and 0.751. These analyses show that discrimination did not depend "
                "entirely on NIV, while the remaining temporal decrement still argues against immediate "
                "transportability. NIV is now interpreted as a treatment/documentation variable and not as "
                "a causal driver of mortality."
            ),
            (
                "The missing-data and sensitivity-analysis Methods, the cohort-specific Results, the "
                "limitations, and Supplementary Table S12 now report NIV availability, the exclusion "
                "analysis, and its non-causal interpretation."
            ),
        ),
        (
            "Reviewer 3 - Major Concern 6",
            r3p[11],
            (
                "We thank the Reviewer for cautioning against overinterpretation of the NH cohort result. "
                "The NH analysis is now explicitly designated exploratory because it contains only 98 "
                "patients and 37 events. Although the reduced model achieved AUROC 0.898, its "
                "calibration-in-the-large was 0.386 and its calibration slope was 2.387, demonstrating that "
                "high discrimination in this small sample did not translate into reliable absolute-risk "
                "estimation.\n\n"
                "We therefore report the NH estimate as a hypothesis-generating observation, avoid comparing "
                "it as evidence of superior geographic transportability, and state that the wide uncertainty, "
                "small cohort size, and material miscalibration preclude clinical inference."
            ),
            (
                "The Abstract, external-evaluation Results, Figure 4 and Figure 7 legends, and Discussion now "
                "label NH as exploratory and report its sample size, event count, and calibration limitations."
            ),
        ),
        (
            "Reviewer 3 - Minor Concern 1",
            r3p[13],
            (
                "We thank the Reviewer for requesting clarification of the competing role of hospital "
                "discharge. The available endpoint is 28-day in-hospital mortality measured from hospital "
                "admission. Patients discharged alive before day 28 are classified as non-events, and deaths "
                "occurring after discharge are not captured. Thus, the analysis does not estimate 28-day "
                "all-cause mortality and may be affected by informative discharge and between-institution "
                "differences in discharge practice."
            ),
            (
                "The outcome-definition Methods and limitations now state the time origin, handling of live "
                "discharge, absence of post-discharge ascertainment, and the resulting risk of informative "
                "discharge bias."
            ),
        ),
        (
            "Reviewer 3 - Minor Concern 2",
            r3p[14],
            (
                "We thank the Reviewer for asking us to distinguish the status of the reduced model from the "
                "locked benchmark. The 8-feature model was defined after reviewing development-data feature "
                "importance and SHAP results; it was not prespecified in the original protocol. We now label "
                "it consistently as a post hoc, implementation-oriented derivative and acknowledge that its "
                "apparent performance may be optimistic because feature reduction was informed by the "
                "development data."
            ),
            (
                "The model-development Methods, Results, Discussion, Figure 7 legend, and Supplementary Table "
                "S15 now distinguish the prespecified 18-predictor benchmark from the post hoc 8-feature "
                "derivative and specify their respective evaluation roles."
            ),
        ),
        (
            "Reviewer 3 - Minor Concern 3",
            r3p[15],
            (
                "We thank the Reviewer for emphasizing the effective sample size available for model "
                "development. The training subset contained 160 endpoint events for 18 locked predictors, "
                "and the ensemble additionally involved algorithm-specific tuning. This event count is "
                "modest relative to the modelling flexibility and makes overfitting a credible explanation "
                "for the internal-to-temporal performance decline.\n\n"
                "We have therefore avoided presenting the random split as definitive validation and identify "
                "temporal evaluation, confidence intervals, and future prospectively specified validation "
                "as more informative tests of generalizability."
            ),
            (
                "The limitations now report the 160 training events and 18 predictors explicitly and link "
                "the modest effective sample size to possible overfitting and temporal performance loss."
            ),
        ),
        (
            "Reviewer 3 - Minor Concern 4",
            r3p[16],
            (
                "We thank the Reviewer for proposing a title that better reflects the design and evidentiary "
                "scope. We adopted a retrospective-evaluation framing and removed wording that could imply "
                "prospective bedside prediction or demonstrated clinical benefit. The revised title also "
                "specifies that the population was selected by documented hypoferremia."
            ),
            (
                f"The title is now “{NEW_TITLE},” and the same retrospective risk-stratification framing "
                "is used in the Abstract, objectives, Discussion, and Conclusions."
            ),
        ),
    ]
    _add_response_heading(doc, "Reviewer 3", reviewer=True)
    for heading, comment, response, changes in r3_items:
        _add_response_heading(doc, heading)
        _add_response(doc, comment, response, changes)

    r4p = _source_paragraphs("reviewer4.docx")
    r4_ranges = [
        ("Reviewer 4 - Overall Assessment", 2, 10),
        ("Reviewer 4 - Major Comment 1", 17, 21),
        ("Reviewer 4 - Major Comment 2", 25, 30),
        ("Reviewer 4 - Major Comment 3", 34, 39),
        ("Reviewer 4 - Major Comment 4", 43, 47),
        ("Reviewer 4 - Major Comment 5", 50, 54),
        ("Reviewer 4 - Major Comment 6", 59, 63),
        ("Reviewer 4 - Major Comment 7", 67, 71),
        ("Reviewer 4 - Major Comment 8", 75, 79),
        ("Reviewer 4 - Major Comment 9", 83, 91),
        ("Reviewer 4 - Moderate Comment 1", 98, 99),
        ("Reviewer 4 - Moderate Comment 2", 103, 104),
        ("Reviewer 4 - Moderate Comment 3", 108, 112),
        ("Reviewer 4 - Moderate Comment 4", 116, 118),
        ("Reviewer 4 - Moderate Comment 5", 122, 125),
        ("Reviewer 4 - Minor Comment 1", 129, 129),
        ("Reviewer 4 - Minor Comment 2", 132, 132),
        ("Reviewer 4 - Minor Comment 3", 135, 135),
        ("Reviewer 4 - Minor Comment 4", 138, 138),
        ("Reviewer 4 - Minor Comment 5", 141, 141),
        ("Reviewer 4 - Minor Comment 6", 144, 145),
        ("Reviewer 4 - Minor Comment 7", 148, 148),
        ("Reviewer 4 - Minor Comment 8", 151, 152),
        ("Reviewer 4 - Summary and Final Recommendation", 155, 163),
    ]
    r4_answers = [
        (
            "We thank the Reviewer for this detailed and constructive reassessment and for recognizing the "
            "substantial changes made after the first review. We agree that the remaining issues concern the "
            "biological meaning of hypoferremia, predictor timing, calibration, transportability, and reporting "
            "completeness. We addressed them through a reproducibility audit, additional prespecified and "
            "sensitivity analyses, and a cross-document review of terminology and claims.\n\n"
            "The audit also identified limitations not fully apparent in the prior draft, including use of "
            "whole-ICU-stay summaries, a corrected MIMIC-III denominator, and inability to reproduce the former "
            "normal-iron comparator. We report these findings transparently and restrict the revised "
            "interpretation to retrospective mortality-risk stratification within patients selected by "
            "documented biochemical hypoferremia.",
            "The manuscript, Supplementary Information, figures, TRIPOD+AI checklist, and reproducibility "
            "materials were revised and cross-checked; the specific changes and their locations are documented "
            "under each comment below.",
        ),
        (
            "We thank the Reviewer for requesting a fuller pathophysiological account of hypoferremia during "
            "critical illness. We expanded the biological rationale to describe how inflammatory signalling, "
            "particularly interleukin-6-mediated hepatic hepcidin induction, promotes ferroportin degradation "
            "and iron sequestration within macrophages and the reticuloendothelial system. In this setting, a "
            "low serum iron concentration may reflect inflammation-associated functional iron restriction and "
            "illness severity; it cannot, by itself, establish depleted total-body iron stores.\n\n"
            "This mechanistic distinction is consistent with the new within-cohort analyses: the adjusted odds "
            "ratio per 10 µg/dL higher first serum iron was 1.052 (95% CI 0.946–1.171; P=0.348), and adding "
            "serum iron to clinical covariates changed out-of-fold AUROC by −0.0014 (95% CI −0.0075 to 0.0044). "
            "Accordingly, the manuscript now treats serum iron primarily as a cohort-defining biomarker and "
            "potential correlate of systemic illness, not a demonstrated causal determinant of mortality.",
            "The serum-iron Results, the mechanistic and limitations paragraphs in the Discussion, and the "
            "Conclusions were revised; Supplementary Tables "
            "S13–S14 provide the supporting adjusted and incremental analyses.",
        ),
        (
            "We thank the Reviewer for highlighting the need for a biologically cautious interpretation of "
            "the former restricted-cubic-spline result. During the audit, we could not reproduce the original "
            "normal-iron comparator because the retained project materials did not include the required "
            "patient-level comparator table or extraction code. We therefore removed that comparison rather "
            "than retain an unverifiable curve.\n\n"
            "In the reproducible analysis restricted to 1,487 eligible patients with documented hypoferremia, "
            "the four-knot restricted-cubic-spline likelihood-ratio test did not support nonlinearity "
            "(P=0.555), and the adjusted linear "
            "association was also not statistically supported. Figure 2 was replaced by the audited within-"
            "cohort adjusted dose-response analysis. The revised Discussion notes that transfusion, liver "
            "injury, tissue damage, malignancy, and multiorgan failure could confound serum iron, but it does "
            "not invoke these mechanisms to explain a pattern that the audited data did not confirm.",
            "Figure 2, its legend, the statistical-analysis Methods, the serum-iron Results, and the Discussion "
            "were replaced or revised; Supplementary Table S13 reports the adjusted effect and nonlinearity test.",
        ),
        (
            "We thank the Reviewer for requesting an explicit predictor observation window, which was crucial "
            "for interpreting the intended use of the model. The audit established that variables ending in "
            "_min, _max, and _mean summarize all recorded values over the full ICU stay, while variables "
            "ending in _1st represent the first recorded value during that ICU stay. The retained data did not "
            "contain authentic serum-iron chart times, and the legacy los_icu_to_t0 field was a synthetic proxy "
            "rather than a verified clinical prediction time.\n\n"
            "Because some summaries may incorporate information recorded after early clinical deterioration, "
            "we no longer present the model as prediction at admission, within 24 hours, or before a defined "
            "bedside decision. It is now described as retrospective whole-stay risk stratification. This "
            "correction resolves the ambiguity but also limits prospective clinical interpretation, which is "
            "stated explicitly.",
            "Methods section 2a now defines every suffix and observation window. Supplementary Table S6, the "
            "TRIPOD+AI checklist, model card, variable dictionary, and repository documentation were synchronized "
            "with the retrospective whole-stay framing.",
        ),
        (
            "We thank the Reviewer for stressing that discrimination and calibration answer complementary "
            "questions. We added Brier score, calibration-in-the-large, and calibration slope with 95% "
            "confidence intervals for the benchmark and reduced models wherever patient-level predictions were "
            "available. For the benchmark, internal Brier score, calibration-in-the-large, and slope were "
            "0.109, −0.898, and 1.705; the corresponding corrected MIMIC-III values were 0.141, −0.597, and "
            "0.929.\n\n"
            "These results show clinically relevant probability miscalibration despite moderate discrimination, "
            "so the manuscript no longer uses AUROC as sufficient evidence for clinical readiness. The small "
            "numbers of events in several evaluation cohorts are also acknowledged as limiting the precision of "
            "calibration estimates and strengthening the requirement for local recalibration.",
            "Calibration definitions were added to Methods; the Results and Table 3 report principal estimates; "
            "Supplementary Table S10 provides the structured cohort summary and confidence intervals; and the "
            "Discussion explains their transportability implications.",
        ),
        (
            "We thank the Reviewer for asking us to clarify how and when the reduced model was developed. The "
            "8-feature model was created after reviewing feature rankings and SHAP explanations in the "
            "development data; it was not prespecified in the original analytical protocol. It is therefore "
            "reported as a post hoc, implementation-oriented derivative, and its performance is subject to "
            "optimism because feature reduction was informed by the same development process.\n\n"
            "The prespecified 18-predictor ensemble remains the protocol-continuity benchmark, whereas the "
            "reduced model is presented as a secondary candidate for future, independently specified "
            "evaluation. We do not claim that it replaces the benchmark or that its smaller feature set has "
            "already established implementation superiority.",
            "The model-development Methods, Results, Discussion, Figure 7 legend, and Supplementary Table S15 "
            "were revised to distinguish prespecified and post hoc status, development pathway, evaluation "
            "cohorts, and intended role.",
        ),
        (
            "We thank the Reviewer for emphasizing the distinction between predictive attribution and causal "
            "explanation. We now describe SHAP values as conditional summaries of how variables contribute to "
            "the fitted model's predictions in the analysed data. A high absolute SHAP value indicates "
            "predictive influence under the fitted model; it does not show that changing the variable would "
            "change mortality risk.\n\n"
            "This qualification is particularly important for treatment and documentation variables such as "
            "ventilation and for inflammatory biomarkers such as serum iron, all of which may be affected by "
            "disease severity, clinical decisions, or measurement processes. We removed mechanistic language "
            "that could imply causal discovery.",
            "The interpretability Methods, SHAP Results, Discussion, and Figure 6 and Figure 7 legends were "
            "revised to use predictive-attribution terminology and to exclude causal interpretation.",
        ),
        (
            "We thank the Reviewer for drawing attention to cross-system differences in missing-data mechanisms. "
            "We now report predictor-specific missingness and acknowledge that an imputer fitted in MIMIC-IV may "
            "encode its measurement frequency, documentation practice, and case mix; applying that object in a "
            "new health system does not ensure that the missing-at-random assumption or imputation relationships "
            "remain valid.\n\n"
            "This concern was most visible for NIV, which was missing in 468/799 (58.6%) corrected MIMIC-III "
            "admissions. We therefore added an NIV-exclusion analysis: the 17-feature benchmark produced AUROC "
            "0.816 internally and 0.727 in corrected MIMIC-III, while the 7-feature derivative produced 0.783 "
            "and 0.751. The sensitivity analysis reduces concern about complete dependence on NIV but does not "
            "remove the broader missingness-transportability limitation.",
            "The missing-data Methods, cohort-specific Results, limitations, Supplementary Table S6, and "
            "Supplementary Table S12 were revised to report missingness, imputation transportability, and the "
            "NIV-exclusion results.",
        ),
        (
            "We thank the Reviewer for placing the validation results within the broader problem of model "
            "transportability. We now distinguish the corrected MIMIC-III analysis as same-institution temporal "
            "evaluation from the small F3 and NH analyses as exploratory geographic evaluations of the post hoc "
            "reduced model. This separation is important because each design addresses a different source of "
            "distribution shift.\n\n"
            "The Discussion now identifies laboratory methods, documentation density, ICU admission policies, "
            "case mix, treatment practices, and baseline risk as plausible sources of performance and "
            "calibration variation. We state that the present retrospective results do not establish universal "
            "transportability; prospective multicenter evaluation, site-specific recalibration, clinical "
            "governance, and impact analysis remain necessary before routine use.",
            "The Title, Abstract, cohort-design Methods, Results, Discussion, Conclusions, and web-demonstration "
            "text were revised to separate temporal from geographic evaluation and to moderate implementation claims.",
        ),
        (
            "We thank the Reviewer for specifying the remaining TRIPOD+AI reporting domains. We revised the "
            "reporting checklist and the corresponding manuscript sections to document the full-ICU predictor "
            "window, sequential participant flow, outcome time origin, calibration metrics, missing-data "
            "sensitivity, post hoc status of the reduced model, and the distinction between temporal and "
            "geographic evaluation.\n\n"
            "For computational reproducibility, release v1.1.0 includes the prepared-table analysis code, a "
            "retained iron/timing SQL template, locked preprocessing and model objects, aggregate results, data "
            "dictionary, model card, and Zenodo metadata. Patient-level records, labels, and prediction "
            "probabilities remain controlled. The repository does not claim a complete raw-database extraction "
            "pipeline. The Zenodo archive is staged but unpublished; the corresponding author must publish it "
            "and insert the assigned DOI before final resubmission.",
            "The TRIPOD+AI checklist, Methods, Supplementary Tables S6–S15, Code and Data Availability statements, "
            "repository documentation, model card, and release metadata were updated and cross-referenced.",
        ),
        (
            "We thank the Reviewer for the specific recommendation to improve the Discussion's economy and "
            "narrative focus. We reorganized the section around four distinct questions: biological "
            "interpretation, model performance and calibration, transportability and implementation, and "
            "limitations. Repeated statements about retrospective design, future validation, and clinical "
            "readiness were consolidated rather than restated after each result.\n\n"
            "The resulting Discussion is more concise while retaining the new calibration, missingness, "
            "serum-iron, and transportability interpretations required for scientific completeness.",
            "The Discussion was line-edited and reorganized into focused paragraphs, with duplicated "
            "implementation and validation statements removed or consolidated.",
        ),
        (
            "We thank the Reviewer for recommending a broader comparison with contemporary explainable machine-"
            "learning studies in critical care. We revised the literature comparison so that it no longer "
            "centres on AUROC alone. The Discussion now contrasts feature engineering, the conditional nature "
            "of SHAP explanations, calibration reporting, temporal performance degradation, external "
            "evaluation design, and the practical burden of model maintenance.\n\n"
            "This comparison positions the study as an example of audited retrospective risk stratification "
            "rather than as evidence that explainability or algorithmic complexity is itself sufficient for "
            "clinical translation.",
            "The comparative Discussion was revised to address explainability, calibration, validation design, "
            "and model maintenance rather than relying on AUROC comparisons alone.",
        ),
        (
            "We thank the Reviewer for this clinically important observation about the composition of the "
            "reduced model. The retained features predominantly represent respiratory support, renal "
            "dysfunction, haematologic status, and general physiologic or documentation-related severity; "
            "they do not directly measure iron metabolism. This pattern is concordant with the adjusted and "
            "incremental analyses showing little independent predictive contribution from first serum iron.\n\n"
            "We therefore interpret the model as estimating mortality risk among patients selected by "
            "documented hypoferremia, not as an iron-driven model or a mechanistic model of iron biology. This "
            "distinction is now explicit in the main interpretation.",
            "The reduced-model Results, biological interpretation in the Discussion, and Conclusions were "
            "revised; Supplementary Tables S13–S15 provide the quantitative and developmental context.",
        ),
        (
            "We thank the Reviewer for requesting a clinically accessible explanation of decision-curve "
            "thresholds. We now define the threshold probability as the predicted risk at which the expected "
            "benefit of acting on a true-positive prediction is considered sufficient to justify the relative "
            "harm of a false positive. The threshold is therefore decision-specific and cannot be selected "
            "solely from the model's statistical output.\n\n"
            "Because the study did not prespecify a treatment, consultation, monitoring action, or acceptable "
            "harm-benefit trade-off, no threshold range is claimed to be clinically optimal. Decision-curve "
            "analysis is presented as exploratory and cannot establish net clinical benefit without a defined "
            "use case and prospective impact evaluation.",
            "The Results and Figure 5 legend now define threshold probability, identify the absence of a "
            "prespecified clinical action, and qualify decision-curve findings as exploratory.",
        ),
        (
            "We thank the Reviewer for identifying residual wording that could still be read as implying "
            "readiness for bedside use. We completed a manuscript-wide claim audit and now describe both "
            "models as retrospective research tools. The public interface is labelled a research demonstration "
            "and is not presented as a regulated medical device or clinical decision-support system.\n\n"
            "Every implementation statement now specifies the outstanding requirements: prospective "
            "multicenter validation, local recalibration, workflow evaluation, governance and "
            "regulatory review, and impact analysis tied to a defined clinical action.",
            "The Abstract, Discussion, Conclusions, research-demonstration section, model card, and relevant "
            "figure legends were revised to use consistent pre-deployment qualifications.",
        ),
        (
            "We thank the Reviewer for this precise editorial observation. We reviewed paragraph openings "
            "throughout the Abstract, Introduction, Methods, Results, and Discussion and removed repetitive "
            "uses of “Furthermore,” “Moreover,” and “In addition.” Where a transition remained necessary, it "
            "was replaced with a content-specific connection rather than a generic additive phrase.",
            "The full manuscript was line-edited for varied, content-driven transitions and smoother narrative flow.",
        ),
        (
            "We thank the Reviewer for noting that several long sentences reduced accessibility. We divided "
            "multi-clause sentences where independent methodological or interpretive points had been combined, "
            "while preserving numerical qualifiers and causal limitations. Particular attention was given to "
            "the Abstract, predictor-timing Methods, calibration Results, and limitations.",
            "The Abstract, Methods, Results, and Discussion were line-edited to shorten complex sentences and "
            "improve readability for an international clinical audience.",
        ),
        (
            "We thank the Reviewer for identifying repetition of the phrase “future validation.” We consolidated "
            "overlapping statements into one structured implementation paragraph and used more precise terms "
            "elsewhere—temporal evaluation, geographic evaluation, prospective multicenter validation, local "
            "recalibration, and impact assessment—according to the specific methodological requirement.",
            "The Discussion was shortened and reorganized, and repeated generic references to future validation "
            "were replaced by specific, non-duplicative requirements.",
        ),
        (
            "We thank the Reviewer for requesting uniform biological terminology. We standardized "
            "“documented biochemical hypoferremia” when defining the study population and used “hypoferremia” "
            "only as an unambiguous abbreviated form after definition. We removed remaining instances that "
            "equated serum iron <60 µg/dL with proven iron deficiency or depleted iron stores.",
            "Terminology was audited across the title, Abstract, main text, figure legends, three-line tables, "
            "Supplementary Information, TRIPOD+AI checklist, and repository documentation.",
        ),
        (
            "We thank the Reviewer for noting that the previous flow diagram was visually dense. We retained "
            "the updated author-supplied BioRender Figure 1 requested for this revision, with Panel A providing "
            "a concise conceptual workflow and Panel B separating cohort sources, exclusions, development, "
            "temporal evaluation, and exploratory geographic evaluation. The caption now carries methodological "
            "details that need not be repeated within the artwork.\n\n"
            "The newest image displays the corrected MIMIC-III temporal evaluation cohort as n=799, matching "
            "all Round 2 analyses. The legend also reports the sequential MIMIC-IV screening denominators "
            "that are not separately drawn within the artwork.",
            "Figure 1 was replaced with the updated BioRender export, and its legend was revised for provenance, "
            "denominator transparency, and clearer separation of conceptual and cohort-flow information.",
        ),
        (
            "We thank the Reviewer for suggesting a compact workflow diagram. Panel A of the updated Figure 1 "
            "now summarizes data collection, cohort assembly, feature extraction, machine learning, internal "
            "and external evaluation, and the web-based research demonstration. Panel B provides the detailed "
            "cohort flow, allowing the conceptual pipeline and participant accounting to be read separately.",
            "Figure 1A serves as the simplified workflow panel, while Figure 1B retains the cohort-specific "
            "information required for transparent reporting.",
        ),
        (
            "We thank the Reviewer for requesting explicit clarification of the interface's regulatory and "
            "clinical status. The interface is now described as a research demonstration that illustrates "
            "model input and output behaviour. It is not approved for diagnosis, triage, treatment selection, "
            "or other clinical decision-making, and its displayed probabilities should not be interpreted as "
            "locally calibrated bedside risks. We also corrected the current repository interface so that its "
            "output is consistently labelled as 28-day mortality probability. Because the manuscript image is "
            "an archived screenshot from before that label correction, its legend now states this explicitly "
            "and explains that the displayed survival probability is the complementary probability.",
            "The interface section heading and text, Figure 8 legend, model card, and repository README were "
            "revised to state the research-only purpose, absence of regulatory or clinical approval, and the "
            "relationship between the archived screenshot label and the corrected current interface.",
        ),
        (
            "We thank the Reviewer for recommending a concise comparison of the two model pathways. We added "
            "a dedicated table that identifies the 18-predictor model as the prespecified locked benchmark and "
            "the 8-feature model as a post hoc derivative. The table also distinguishes predictor count, "
            "development basis, evaluation cohorts, intended research role, and the principal caution attached "
            "to each model.\n\n"
            "To ensure consistency with the rest of the submission, this new table uses the same three-line "
            "format as every other statistical table in the manuscript and Supplementary Information.",
            "Supplementary Table S15 was added and formatted as a three-line table; the Methods and Discussion "
            "cross-reference its distinction between benchmark and derivative models.",
        ),
        (
            "We thank the Reviewer for the encouraging final assessment and recommendation. We have addressed "
            "each clarification individually and used the review to strengthen biological interpretation, "
            "predictor-timing transparency, calibration reporting, missing-data analysis, transportability "
            "language, and computational reproducibility. Where the audit did not support a prior result—most "
            "notably the normal-iron comparator—we removed the claim rather than preserve an unverifiable "
            "analysis.\n\n"
            "We also completed a cross-document consistency check so that the title, cohort denominators, "
            "endpoint definition, model status, numerical results, limitations, tables, figures, checklist, "
            "and public release documentation express the same evidentiary scope.",
            "The corresponding revisions are identified under each individual comment above.",
        ),
    ]
    _add_response_heading(doc, "Reviewer 4", reviewer=True)
    for (heading, start, end), (response, changes) in zip(r4_ranges, r4_answers):
        _add_response_heading(doc, heading)
        _add_response(doc, _join_inclusive(r4p, start, end), response, changes)

    r5p = _source_paragraphs("reviewer5.docx")
    _add_response_heading(doc, "Reviewer 5", reviewer=True)
    _add_response_heading(doc, "Reviewer 5 - Comment")
    _add_response(
        doc,
        r5p[1],
        (
            "We thank the Reviewer for identifying the central inconsistency between the former iron-centred "
            "framing and the actual predictors and results. We addressed this concern by testing the adjusted "
            "association, functional form, and incremental predictive contribution of first serum iron in all "
            "1,487 eligible MIMIC-IV admissions. The adjusted odds ratio per 10 µg/dL higher iron was 1.052 "
            "(95% CI 0.946–1.171; P=0.348); the complete-case sensitivity was concordant (OR 1.055; "
            "95% CI 0.930–1.198; P=0.406). The four-knot restricted-cubic-spline likelihood-ratio test "
            "was P=0.555. Adding iron to the clinical covariates changed five-fold out-of-fold AUROC by "
            "−0.0014 (95% CI −0.0075 to 0.0044) and Brier score by 0.0004 (95% CI −0.00005 to "
            "0.00077). These analyses do not support an independent or clinically meaningful incremental "
            "predictive contribution of serum iron within this selected cohort.\n\n"
            "The former normal-iron comparator could not be reproduced because the retained materials lacked "
            "both patient-level comparator records and the extraction code needed to reconstruct them. We "
            "therefore removed the analysis and replaced Figure 2 with the auditable within-hypoferremia "
            "dose-response result. The title, objectives, interpretation, and conclusions now present the "
            "models as general severity-risk stratification among ICU patients selected by documented "
            "biochemical hypoferremia—not as iron-driven models and not as evidence that serum iron causally "
            "determines mortality."
        ),
        (
            "Figure 2 and its legend were replaced; Supplementary Tables S13–S14 were added in three-line "
            "format; and the Title, Abstract, Methods, Results, Discussion, and Conclusions were revised to "
            "match the audited serum-iron findings and the narrower scientific claim."
        ),
    )
    out = R2 / "Response_to_Reviewers_R2.docx"
    doc.save(out)
    return out


def build_editor_letter() -> Path:
    doc = _blank_document_from_template(R1 / "Letter_to_Editor_R1.docx")
    _set_document_metadata(doc, f"Revision Submission - Round 2 — {NEW_TITLE}", "Round 2 cover letter")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Revision Submission - Round 2")
    r.bold = True
    _set_times_new_roman(r, size_pt=16)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _set_times_new_roman(
        p.add_run("Submission ID: 9974989f-eff3-4bd0-a4dc-a3d1280eb44d"),
        size_pt=11,
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _set_times_new_roman(p.add_run(f"Manuscript Title: {NEW_TITLE}"), size_pt=11)
    body_paragraphs = [
        doc.add_paragraph("Dear Dr Secgin and the Editorial Team,"),
        doc.add_paragraph(
        f"We are pleased to resubmit “{NEW_TITLE}” (Submission ID "
        "9974989f-eff3-4bd0-a4dc-a3d1280eb44d). Before revision, we completed a "
        "code-and-data audit and additional statistical analyses. We corrected cohort flow and the "
        "MIMIC-III eligibility denominator, reported standard calibration metrics, added multiple-"
        "imputation and NIV-exclusion sensitivities, and quantified the adjusted and incremental role "
        "of serum iron. The results led us to narrow the manuscript to retrospective severity-risk "
        "stratification among patients selected for documented hypoferremia."
        ),
        doc.add_paragraph(
        "We also removed the previous normal-iron comparator analysis because the retained project "
        "materials did not contain a reproducible comparator dataset or extraction code. Figure 1 was "
        "replaced with the updated author-prepared BioRender export supplied for this revision. Its legend "
        "states that it was created by the authors with BioRender.com. The artwork now displays the corrected "
        "MIMIC-III temporal evaluation cohort (n=799), consistent with every current analysis."
        ),
        doc.add_paragraph(
        "Repository release version 1.1.0 has been prepared for DOI archiving with prepared-table analysis "
        "code, model objects, aggregate R2 results, citation metadata, and release documentation. Patient-level MIMIC "
        "records, labels, and probabilities are excluded from public redistribution under the applicable "
        "data-use conditions. The Zenodo archive has not yet been published; therefore, the revision does "
        "not state or imply that a DOI already exists. After author verification, the archive must be published "
        "and its assigned DOI inserted before final resubmission."
        ),
        doc.add_paragraph(
        "We believe these changes directly address the editor’s and reviewers’ concerns while making the "
        "study’s scope, limitations, and reproducibility substantially clearer."
        ),
        doc.add_paragraph("Sincerely,"),
        doc.add_paragraph("Yang Xi, on behalf of all authors"),
    ]
    for paragraph in body_paragraphs:
        for run in paragraph.runs:
            _set_times_new_roman(run, size_pt=11)
    out = R2 / "Letter_to_Editor_R2.docx"
    doc.save(out)
    return out


def copy_release_materials() -> None:
    for obsolete_name in [
        "figure1_study_design.png",
        "figure1_study_design.svg",
        "figure1_study_design.pdf",
        "plot_figure1_study_design.py",
        "figure1_original_biorender.png",
        "figure1_original_biorender.tiff",
    ]:
        obsolete = R2 / obsolete_name
        if obsolete.exists():
            obsolete.unlink()
    for name in [
        ".zenodo.json",
        "CITATION.cff",
        "CITATION.bib",
        "CODE_AVAILABILITY.md",
        "DATA_AVAILABILITY.md",
        "MODEL_CARD.md",
        "RELEASE_NOTES_v1.1.0.md",
        "RELEASE_CHECKLIST.md",
    ]:
        shutil.copy2(ROOT / name, R2 / name)
    shutil.copy2(REF / "FIGURE1_SOURCE_DECLARATION.md", R2 / "FIGURE1_SOURCE_DECLARATION.md")
    with Image.open(UPDATED_FIGURE1) as im:
        im.convert("RGB").save(
            REF / "figure1_biorender_updated.tiff",
            format="TIFF",
            compression="tiff_lzw",
            dpi=(300, 300),
        )
    shutil.copy2(UPDATED_FIGURE1, R2 / "figure1_biorender_updated.png")
    shutil.copy2(REF / "figure1_biorender_updated.tiff", R2 / "figure1_biorender_updated.tiff")
    for name in [
        "figure2_within_hypoferremia_dose_response.png",
        "figure2_within_hypoferremia_dose_response.svg",
        "figure4_corrected_cohort_performance.png",
        "figure4_corrected_cohort_performance.svg",
        "figure7_reduced_model_corrected.png",
    ]:
        shutil.copy2(REF / name, R2 / name)
    shutil.copy2(ROOT / "scripts" / "plot_r2_performance_figures.py", R2 / "plot_r2_performance_figures.py")
    shutil.copy2(ROOT / "scripts" / "r2_supplementary_analyses.py", R2 / "r2_supplementary_analyses.py")


def main() -> None:
    R2.mkdir(parents=True, exist_ok=True)
    build_manuscript(False)
    build_manuscript(True)
    build_supplement(False)
    build_supplement(True)
    build_tripod()
    build_response_letter()
    build_editor_letter()
    copy_release_materials()
    print(f"R2 package written to {R2}")


if __name__ == "__main__":
    main()
