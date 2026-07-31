#!/usr/bin/env python3
"""Structural and cross-document QA for the Round 2 submission package."""
from __future__ import annotations

import hashlib
import json
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Pt

from build_r2_revision_package import MANUSCRIPT_REPLACEMENTS


ROOT = Path(__file__).resolve().parent.parent
PACKAGE = ROOT / "deliverables" / "Fe_R2_Revision_Package"
R1 = ROOT / ".cursor" / "low_fe_3"
SOURCES = R1 / "r2_review_sources"
FIGURE1 = ROOT / "reference" / "r2" / "figure1_biorender_updated.png"


def _text_payload(doc: Document) -> dict:
    return {
        "paragraphs": [paragraph.text for paragraph in doc.paragraphs],
        "tables": [
            [[cell.text for cell in row.cells] for row in table.rows]
            for table in doc.tables
        ],
    }


def _yellow_paragraphs(doc: Document) -> list[str]:
    return [
        paragraph.text
        for paragraph in doc.paragraphs
        if any(run.font.highlight_color == WD_COLOR_INDEX.YELLOW for run in paragraph.runs)
    ]


def _assert_page_geometry_matches(
    reference: Document,
    revised: Document,
    label: str,
) -> None:
    assert len(reference.sections) == len(revised.sections), (
        f"{label}: section count changed from {len(reference.sections)} "
        f"to {len(revised.sections)}"
    )
    fields = (
        "page_width",
        "page_height",
        "top_margin",
        "bottom_margin",
        "left_margin",
        "right_margin",
        "header_distance",
        "footer_distance",
    )
    for section_index, (old, new) in enumerate(
        zip(reference.sections, revised.sections)
    ):
        for field in fields:
            assert getattr(old, field) == getattr(new, field), (
                f"{label}: section {section_index} {field} changed "
                f"from {getattr(old, field)} to {getattr(new, field)}"
            )


def _assert_explicit_tnr(run, label: str) -> None:
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    assert run.font.name == "Times New Roman", (
        f"{label}: highlighted run does not explicitly use Times New Roman"
    )
    for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        assert rfonts.get(qn(f"w:{slot}")) == "Times New Roman", (
            f"{label}: highlighted run has non-Times-New-Roman {slot} slot"
        )


def _assert_highlighted_typography(doc: Document, label: str) -> int:
    checked = 0
    for paragraph_index, paragraph in enumerate(doc.paragraphs):
        for run_index, run in enumerate(paragraph.runs):
            if run.text and run.font.highlight_color == WD_COLOR_INDEX.YELLOW:
                _assert_explicit_tnr(
                    run,
                    f"{label}: paragraph {paragraph_index}, run {run_index}",
                )
                assert run.font.size is not None, (
                    f"{label}: highlighted paragraph {paragraph_index}, "
                    f"run {run_index} has no explicit font size"
                )
                checked += 1
    for table_index, table in enumerate(doc.tables):
        for row_index, row in enumerate(table.rows):
            for cell_index, cell in enumerate(row.cells):
                for paragraph in cell.paragraphs:
                    for run_index, run in enumerate(paragraph.runs):
                        if (
                            run.text
                            and run.font.highlight_color
                            == WD_COLOR_INDEX.YELLOW
                        ):
                            _assert_explicit_tnr(
                                run,
                                (
                                    f"{label}: table {table_index}, row {row_index}, "
                                    f"cell {cell_index}, run {run_index}"
                                ),
                            )
                            assert run.font.size is not None, (
                                f"{label}: highlighted table run has no explicit size"
                            )
                            checked += 1
    assert checked, f"{label}: no highlighted text runs were checked"
    return checked


def _assert_response_typography(path: Path) -> None:
    doc = Document(path)
    allowed_sizes = {11.0, 12.0, 16.0}
    for paragraph_index, paragraph in enumerate(doc.paragraphs):
        for run_index, run in enumerate(paragraph.runs):
            if not run.text:
                continue
            _assert_explicit_tnr(
                run,
                f"{path.name}: paragraph {paragraph_index}, run {run_index}",
            )
            assert run.font.size is not None
            assert round(run.font.size.pt, 1) in allowed_sizes, (
                f"{path.name}: unexpected {run.font.size.pt:g}-point text"
            )


def _first_text_run(paragraph):
    return next((run for run in paragraph.runs if run.text), None)


def _assert_replacement_hierarchy_matches(
    reference: Document,
    revised: Document,
    paragraph_indices: tuple[int | tuple[int, int], ...],
    label: str,
) -> None:
    for item in paragraph_indices:
        reference_index, revised_index = item if isinstance(item, tuple) else (item, item)
        old_run = _first_text_run(reference.paragraphs[reference_index])
        new_run = _first_text_run(revised.paragraphs[revised_index])
        assert old_run is not None and new_run is not None
        old_size = (
            old_run.font.size
            or reference.paragraphs[reference_index].style.font.size
            or Pt(12)
        )
        assert new_run.font.size == old_size, (
            f"{label}: paragraph {revised_index} size changed from "
            f"{old_size.pt:g} pt to "
            f"{new_run.font.size.pt if new_run.font.size else 'inherited'}"
        )
        assert bool(new_run.bold) == bool(old_run.bold), (
            f"{label}: paragraph {revised_index} bold hierarchy changed"
        )
        assert (
            revised.paragraphs[revised_index].paragraph_format.space_before
            == reference.paragraphs[reference_index].paragraph_format.space_before
        ), f"{label}: paragraph {revised_index} space-before changed"
        assert (
            revised.paragraphs[revised_index].paragraph_format.space_after
            == reference.paragraphs[reference_index].paragraph_format.space_after
        ), f"{label}: paragraph {revised_index} space-after changed"


def _assert_three_line_tables(doc: Document, label: str) -> None:
    for table_index, table in enumerate(doc.tables):
        if not table.rows:
            continue
        for row_index, row in enumerate(table.rows):
            for cell_index, cell in enumerate(row.cells):
                borders = cell._tc.get_or_add_tcPr().find(qn("w:tcBorders"))
                assert borders is not None, (
                    f"{label}: table {table_index}, cell {row_index}/{cell_index} "
                    "has no explicit borders"
                )
                for edge in ("start", "end", "insideH", "insideV"):
                    node = borders.find(qn(f"w:{edge}"))
                    assert node is not None and node.get(qn("w:val")) == "nil", (
                        f"{label}: table {table_index}, cell {row_index}/{cell_index} "
                        f"contains a non-three-line {edge} border"
                    )
                top = borders.find(qn("w:top"))
                bottom = borders.find(qn("w:bottom"))
                assert top is not None and bottom is not None
                expected_top = "single" if row_index == 0 else "nil"
                expected_bottom = (
                    "single"
                    if row_index in {0, len(table.rows) - 1}
                    else "nil"
                )
                assert top.get(qn("w:val")) == expected_top
                assert bottom.get(qn("w:val")) == expected_bottom


def _assert_no_revision_markup(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml")
        for forbidden in (
            rb"<w:ins(?:\s|>)",
            rb"<w:del(?:\s|>)",
            rb"<w:commentRangeStart(?:\s|>)",
        ):
            assert re.search(forbidden, xml) is None, (
                f"{path.name}: tracked/comment markup found"
            )


def _expected_reviewer_paragraphs() -> list[str]:
    def paragraphs(name: str) -> list[str]:
        return [paragraph.text for paragraph in Document(SOURCES / name).paragraphs]

    editor = paragraphs("editor_and_portal.docx")[1]
    expected: list[str] = []
    blocks = [
        editor[editor.index("Please note that") : editor.index("\n\nEditor Comments")],
        editor[
            editor.index('"The manuscript is generally well written')
            : editor.index("\n-Yusuf Secgin")
        ],
        editor[
            editor.index("Please could you clarify where the images")
            : editor.index("\n\nWe recommend submitting")
        ],
    ]
    for block in blocks:
        expected.extend(part for part in block.split("\n\n") if part)

    r3 = paragraphs("reviewer3.docx")
    expected.extend([r3[3], r3[4], *r3[6:12], *r3[13:17]])

    r4 = paragraphs("reviewer4.docx")
    ranges = [
        (2, 10),
        (17, 21),
        (25, 30),
        (34, 39),
        (43, 47),
        (50, 54),
        (59, 63),
        (67, 71),
        (75, 79),
        (83, 91),
        (98, 99),
        (103, 104),
        (108, 112),
        (116, 118),
        (122, 125),
        (129, 129),
        (132, 132),
        (135, 135),
        (138, 138),
        (141, 141),
        (144, 145),
        (148, 148),
        (151, 152),
        (155, 163),
    ]
    for start, end in ranges:
        expected.extend(text for text in r4[start : end + 1] if text)

    r5 = paragraphs("reviewer5.docx")
    expected.extend(part for part in r5[1].split("\n\n") if part)
    return expected


def _verify_response_letter(path: Path) -> dict:
    doc = Document(path)
    paragraphs = doc.paragraphs
    response_paragraphs = [
        paragraph for paragraph in paragraphs if paragraph.text.startswith("Response: ")
    ]
    location_paragraphs = [
        paragraph
        for paragraph in paragraphs
        if paragraph.text.startswith(
            "Manuscript/Supplementary revision (highlighted in yellow): "
        )
    ]
    assert len(response_paragraphs) == 39, (
        f"Expected 39 point-by-point responses; found {len(response_paragraphs)}"
    )
    assert len(location_paragraphs) == 39, (
        f"Expected 39 change-location statements; found {len(location_paragraphs)}"
    )
    for index, paragraph in enumerate(response_paragraphs, start=1):
        response = paragraph.text.removeprefix("Response: ").lstrip()
        assert response.startswith(("We thank", "We sincerely thank")), (
            f"Response {index} does not begin with thanks"
        )
        assert all(
            run.font.color.rgb is None or str(run.font.color.rgb) == "0070C0"
            for run in paragraph.runs
            if run.text
        ), f"Response {index} is not uniformly blue"
        for run in paragraph.runs:
            if run.text:
                _assert_explicit_tnr(run, f"Response {index}")
                assert run.font.size == Pt(11), (
                    f"Response {index} is not 11-point Times New Roman"
                )
    for index, paragraph in enumerate(location_paragraphs, start=1):
        assert all(
            run.font.color.rgb is None or str(run.font.color.rgb) == "0070C0"
            for run in paragraph.runs
            if run.text
        ), f"Location statement {index} is not uniformly blue"

    response_by_text = {}
    for paragraph in paragraphs:
        response_by_text.setdefault(paragraph.text, []).append(paragraph)
    missing = []
    not_italic_black = []
    for expected in _expected_reviewer_paragraphs():
        candidates = response_by_text.get(expected, [])
        if not candidates:
            missing.append(expected[:100])
            continue
        valid = False
        for paragraph in candidates:
            nonempty = [run for run in paragraph.runs if run.text]
            if nonempty and all(run.italic for run in nonempty) and all(
                run.font.color.rgb is None or str(run.font.color.rgb) == "000000"
                for run in nonempty
            ):
                valid = True
                break
        if not valid:
            not_italic_black.append(expected[:100])
    assert not missing, f"Reviewer text not reproduced verbatim: {missing[:3]}"
    assert not not_italic_black, (
        f"Reviewer text not uniformly italic black: {not_italic_black[:3]}"
    )
    for paragraph in paragraphs:
        nonempty = [run for run in paragraph.runs if run.text]
        if nonempty and all(run.italic for run in nonempty):
            for run in nonempty:
                _assert_explicit_tnr(run, "Reviewer/editor quotation")
                assert run.font.size == Pt(11)
    assert "Summary of Changes" not in "\n".join(p.text for p in paragraphs)
    return {
        "responses": len(response_paragraphs),
        "locations": len(location_paragraphs),
        "verbatim_source_paragraphs": len(_expected_reviewer_paragraphs()),
    }


def _verify_highlight_pair(clean_path: Path, highlighted_path: Path, label: str) -> dict:
    clean = Document(clean_path)
    highlighted = Document(highlighted_path)
    assert _text_payload(clean) == _text_payload(highlighted), (
        f"{label}: clean and highlighted text/table content differ"
    )
    yellow = _yellow_paragraphs(highlighted)
    assert yellow, f"{label}: no highlighted revisions found"
    clean_yellow = _yellow_paragraphs(clean)
    assert not clean_yellow, f"{label}: clean file contains yellow highlighting"
    _assert_three_line_tables(clean, f"{label} clean")
    _assert_three_line_tables(highlighted, f"{label} highlighted")
    highlighted_runs = _assert_highlighted_typography(
        highlighted,
        f"{label} highlighted",
    )
    return {
        "paragraphs": len(clean.paragraphs),
        "tables": len(clean.tables),
        "highlighted_paragraphs": len(yellow),
        "highlighted_runs_with_explicit_tnr": highlighted_runs,
    }


def _verify_manuscript_content(path: Path) -> None:
    doc = Document(path)
    full_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    required = [
        "corrected MIMIC-III temporal cohort (n=799)",
        "0.705 in the corrected MIMIC-III",
        "four-knot restricted cubic spline",
        "P=0.555",
        "ΔAUROC −0.0014",
        "Figure created by the authors with BioRender.com.",
        "retrospective whole-stay risk stratification",
        "does not constitute a complete raw-database-to-modeling-table extraction pipeline",
        "AUROC of the locked 18-predictor benchmark",
    ]
    for phrase in required:
        assert phrase in full_text, f"Manuscript missing required revision: {phrase}"
    forbidden = [
        "0.707 in the corrected MIMIC-III",
        "P=0.711",
        "ΔAUROC −0.0002",
        "Baseline characteristics of all eligible",
        "source/export MIMIC-III denominator",
        "low-iron-selected ICU population",
        "The earlier n=1,046 MIMIC-III results",
    ]
    for phrase in forbidden:
        assert phrase not in full_text, f"Manuscript contains superseded text: {phrase}"

    assert len(doc.inline_shapes) == 8, (
        f"Manuscript must contain all 8 R1 figures; found {len(doc.inline_shapes)}"
    )
    with zipfile.ZipFile(path) as archive:
        image = archive.read("word/media/image1.png")
        assert hashlib.sha256(image).digest() == hashlib.sha256(
            FIGURE1.read_bytes()
        ).digest(), "Embedded Figure 1 is not the newest supplied BioRender export"
        document_xml = archive.read("word/document.xml").decode("utf-8")
        for figure_number in range(1, 9):
            assert f'title="Figure {figure_number}"' in document_xml, (
                f"Embedded Figure {figure_number} is missing or lacks title metadata"
            )
        assert 'descr="Study figure;' in document_xml


def _verify_supplement_content(path: Path) -> None:
    doc = Document(path)
    checklist_text = "\n".join(
        cell.text for row in doc.tables[4].rows for cell in row.cells
    )
    for stale_location in ("Methods 2.1", "Methods 2.2", "Methods 2.3"):
        assert stale_location not in checklist_text, (
            f"Supplementary TRIPOD table contains stale location: {stale_location}"
        )

    imbalance_text = "\n".join(
        cell.text for row in doc.tables[7].rows for cell in row.cells
    )
    assert "1,046" not in imbalance_text, (
        "Supplementary Table S8 contains superseded MIMIC-III n=1,046 results"
    )
    assert "MIMIC-III" not in imbalance_text, (
        "Supplementary Table S8 should be limited to auditable MIMIC-IV internal results"
    )
    expected_headers = {
        5: ["Variable", "n", "Missing, n", "Missing, %", "Observation window"],
        6: [
            "Step",
            "n",
            "Hospital deaths, n",
            "Hospital mortality, %",
            "28-day in-hospital deaths, n",
            "28-day in-hospital mortality, %",
        ],
        10: [
            "Characteristic",
            "Included complete-case",
            "Excluded for missingness",
            "Absolute SMD",
            "P value",
        ],
    }
    for table_index, expected in expected_headers.items():
        actual = [cell.text for cell in doc.tables[table_index].rows[0].cells]
        assert actual == expected, (
            f"Supplementary table {table_index} has non-publication headers: "
            f"{actual}"
        )


def main() -> int:
    manuscript_clean = PACKAGE / "Hypoferremia_Manuscript_R2_Clean.docx"
    manuscript_highlighted = (
        PACKAGE / "Hypoferremia_Manuscript_R2_Highlighted.docx"
    )
    supplement_clean = PACKAGE / "Hypoferremia_Supplementary_R2_Clean.docx"
    supplement_highlighted = (
        PACKAGE / "Hypoferremia_Supplementary_R2_Highlighted.docx"
    )
    response = PACKAGE / "Response_to_Reviewers_R2.docx"
    editor_letter = PACKAGE / "Letter_to_Editor_R2.docx"
    checklist = PACKAGE / "Hypoferremia_TRIPOD_AI_Checklist_R2.docx"
    manuscript_reference = Document(
        R1 / "Hypoferremia_Manuscript_R1_Clean.docx"
    )
    manuscript_revised = Document(manuscript_highlighted)
    supplement_reference = Document(
        R1 / "Hypoferremia_Supplementary_R1_Clean.docx"
    )
    supplement_revised = Document(supplement_highlighted)

    results = {
        "manuscript_pair": _verify_highlight_pair(
            manuscript_clean, manuscript_highlighted, "manuscript"
        ),
        "supplement_pair": _verify_highlight_pair(
            supplement_clean, supplement_highlighted, "supplement"
        ),
        "response_letter": _verify_response_letter(response),
    }
    _assert_page_geometry_matches(
        manuscript_reference,
        manuscript_revised,
        "manuscript",
    )
    _assert_page_geometry_matches(
        supplement_reference,
        supplement_revised,
        "supplement",
    )
    _assert_replacement_hierarchy_matches(
        manuscript_reference,
        manuscript_revised,
        tuple(
            index for index in MANUSCRIPT_REPLACEMENTS if index not in {62, 63}
        )
        + ((46, 62), (56, 63), 197),
        "manuscript",
    )
    _assert_replacement_hierarchy_matches(
        supplement_reference,
        supplement_revised,
        (0, 1, 4, 5, 7, 10, 13, 19, 20, 21, 22, 23, 24, 25, 26, 27),
        "supplement",
    )
    _verify_manuscript_content(manuscript_highlighted)
    _verify_supplement_content(supplement_highlighted)
    _assert_response_typography(response)
    _assert_response_typography(editor_letter)
    _assert_page_geometry_matches(
        Document(R1 / "Response_to_Reviewers_R1.docx"),
        Document(response),
        "response letter",
    )
    _assert_page_geometry_matches(
        Document(R1 / "Letter_to_Editor_R1.docx"),
        Document(editor_letter),
        "editor letter",
    )
    _assert_three_line_tables(Document(checklist), "TRIPOD+AI checklist")

    for path in PACKAGE.glob("*.docx"):
        if ".tmp." in path.name:
            continue
        _assert_no_revision_markup(path)

    manifest = json.loads(
        (ROOT / "reference" / "r2" / "r2_analysis_manifest.json").read_text(
            encoding="utf-8"
        )
    )
    assert manifest["mimic3_corrected_n"] == 799
    assert manifest["mimic3_corrected_events"] == 140

    report = [
        "# Round 2 submission structural QA",
        "",
        "- Status: PASS",
        f"- Point-by-point responses: {results['response_letter']['responses']}",
        f"- Verbatim reviewer/editor source paragraphs verified: {results['response_letter']['verbatim_source_paragraphs']}",
        f"- Manuscript tables checked as three-line tables: {results['manuscript_pair']['tables']}",
        f"- Supplementary tables checked as three-line tables: {results['supplement_pair']['tables']}",
        f"- Highlighted manuscript paragraphs: {results['manuscript_pair']['highlighted_paragraphs']}",
        f"- Highlighted supplementary paragraphs: {results['supplement_pair']['highlighted_paragraphs']}",
        f"- Highlighted manuscript runs with explicit Times New Roman: {results['manuscript_pair']['highlighted_runs_with_explicit_tnr']}",
        f"- Highlighted supplementary runs with explicit Times New Roman: {results['supplement_pair']['highlighted_runs_with_explicit_tnr']}",
        "- Clean/highlighted text and table content: identical",
        "- R1 page size and margins: preserved",
        "- R1 paragraph font-size, bold hierarchy, and direct spacing: preserved for revised paragraphs",
        "- Reviewer/editor wording: verbatim, italic, black",
        "- Author responses and change locations: blue; every response opens with thanks",
        "- All 8 manuscript figures: present; newest Figure 1 embedded byte-for-byte; alt text present",
        "- Tracked changes/comments: absent",
        "- Current MIMIC-III denominator: n=799, 140 events",
        "",
        "Visual rendering was audited separately in qa_r2_format_final2/.",
        "",
    ]
    (PACKAGE / "R2_STRUCTURAL_QA_REPORT.md").write_text(
        "\n".join(report),
        encoding="utf-8",
    )
    print(json.dumps(results, indent=2))
    print("R2 structural QA PASS")
    return 0


if __name__ == "__main__":
    sys.exit(main())
