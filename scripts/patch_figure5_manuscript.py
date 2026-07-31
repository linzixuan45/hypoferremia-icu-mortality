#!/usr/bin/env python3
"""Sync Figure 5 image and caption in low_fe_3 manuscripts (FE1 n=269 cohort)."""
from __future__ import annotations

import argparse
import hashlib
import io
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

from docx import Document
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
PKG = ROOT / ".cursor" / "low_fe_3"
FIGURE_PATH = ROOT / "result" / "analysis" / "figures" / "figure5_calibration_dca_cic.png"

MANUSCRIPT_FILES = [
    PKG / "Hypoferremia_Manuscript_R1_Clean.docx",
    PKG / "Hypoferremia_Manuscript_R1_Highlighted.docx",
]

MEDIA_IMAGE5 = "word/media/image5.tiff"
LEGACY_IMAGE5_MD5 = "a2ba665ab9dca4b40fdcabd39c6f24ef"

CAPTION_REPLACEMENTS: list[tuple[str, str]] = [
    (
        "Calibration curves for the 12 models in the validation cohort",
        "Calibration curves for the 12 models in the MIMIC-IV internal test set (n=269)",
    ),
]


def _run_figure5_plot() -> None:
    script = ROOT / "scripts" / "plot_figure5_benchmark.py"
    subprocess.run([sys.executable, str(script), "--verify"], cwd=ROOT, check=True)


def _replace_in_paragraph(paragraph, replacements: list[tuple[str, str]]) -> bool:
    text = paragraph.text
    new_text = text
    for old, new in replacements:
        if old in new_text:
            new_text = new_text.replace(old, new)
    if new_text != text:
        if paragraph.runs:
            paragraph.runs[0].text = new_text
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.text = new_text
        return True
    return False


def _patch_caption_text(path: Path) -> dict[str, int]:
    doc = Document(path)
    stats = {"paragraphs": 0, "table_cells": 0}
    for paragraph in doc.paragraphs:
        if _replace_in_paragraph(paragraph, CAPTION_REPLACEMENTS):
            stats["paragraphs"] += 1
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if _replace_in_paragraph(paragraph, CAPTION_REPLACEMENTS):
                        stats["table_cells"] += 1
    doc.save(path)
    return stats


def _figure_to_tiff_bytes(png_path: Path) -> bytes:
    buf = io.BytesIO()
    Image.open(png_path).convert("RGB").save(buf, format="TIFF", compression="tiff_deflate")
    return buf.getvalue()


def _replace_docx_image(docx_path: Path, png_path: Path) -> str:
    tiff_bytes = _figure_to_tiff_bytes(png_path)
    new_md5 = hashlib.md5(tiff_bytes).hexdigest()
    tmp_path = docx_path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(tmp_path, "w") as zout:
        for info in zin.infolist():
            data = tiff_bytes if info.filename == MEDIA_IMAGE5 else zin.read(info.filename)
            zout.writestr(info, data)
    tmp_path.replace(docx_path)
    return new_md5


def _verify_docx(path: Path, new_image_md5: str) -> list[str]:
    errors: list[str] = []
    with zipfile.ZipFile(path) as z:
        if MEDIA_IMAGE5 not in z.namelist():
            errors.append(f"{path.name}: missing {MEDIA_IMAGE5}")
            return errors
        image_md5 = hashlib.md5(z.read(MEDIA_IMAGE5)).hexdigest()
        xml = z.read("word/document.xml").decode("utf-8", errors="ignore")

    if image_md5 != new_image_md5:
        errors.append(f"{path.name}: embedded image5 MD5 mismatch after patch")
    if image_md5 == LEGACY_IMAGE5_MD5:
        errors.append(f"{path.name}: image5 still legacy MD5 (n≈175 cohort)")

    text = re.sub(r"<[^>]+>", "", xml)
    if "MIMIC-IV internal test set (n=269)" not in text:
        errors.append(f"{path.name}: updated Figure 5 caption with n=269 not found")
    if "Calibration curves for the 12 models in the validation cohort" in text:
        errors.append(f"{path.name}: legacy 'validation cohort' Figure 5 caption remains")

    return errors


def main() -> None:
    parser = argparse.ArgumentParser(description="Patch low_fe_3 Figure 5 from FE1 benchmark")
    parser.add_argument(
        "--regen-figure",
        action="store_true",
        help="Run plot_figure5_benchmark.py --verify first",
    )
    parser.add_argument("--backup", action="store_true", help="Save .bak copies before patching")
    args = parser.parse_args()

    if args.regen_figure:
        _run_figure5_plot()
    elif not FIGURE_PATH.exists():
        raise SystemExit(f"Missing {FIGURE_PATH}. Run with --regen-figure or plot_figure5_benchmark.py first.")

    for manuscript in MANUSCRIPT_FILES:
        if not manuscript.exists():
            raise SystemExit(f"Missing manuscript: {manuscript}")
        if args.backup:
            shutil.copy2(manuscript, manuscript.with_suffix(".docx.bak"))

        stats = _patch_caption_text(manuscript)
        image_md5 = _replace_docx_image(manuscript, FIGURE_PATH)
        errors = _verify_docx(manuscript, image_md5)

        print(f"Patched {manuscript.name}: {stats}, image5_md5={image_md5[:12]}...")
        if errors:
            for err in errors:
                print(f"  VERIFY FAIL: {err}")
            raise SystemExit(1)
        print("  verify OK (Figure 5 image + caption n=269 synced)")


if __name__ == "__main__":
    main()
