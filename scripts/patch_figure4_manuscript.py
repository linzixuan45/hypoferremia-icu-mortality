#!/usr/bin/env python3
"""Sync Figure 4 image and Table 2/3 benchmark AUROCs in low_fe_3 manuscripts."""
from __future__ import annotations

import argparse
import io
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import pandas as pd
from docx import Document
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
PKG = ROOT / ".cursor" / "low_fe_3"
FIGURE_PATH = ROOT / "result" / "analysis" / "figures" / "figure4_roc_12models.png"
TABLE2_CSV = ROOT / "result" / "analysis" / "tables" / "table2_benchmark_12models.csv"
TABLE3_CSV = ROOT / "result" / "analysis" / "tables" / "table3_benchmark_12models.csv"

MANUSCRIPT_FILES = [
    PKG / "Hypoferremia_Manuscript_R1_Clean.docx",
    PKG / "Hypoferremia_Manuscript_R1_Highlighted.docx",
]

MEDIA_IMAGE4 = "word/media/image4.tiff"
INTERNAL_MARKER = "Comparable internal discrimination"
EXTERNAL_MARKER = "Comparable external discrimination"


def _run_benchmark_plot() -> None:
    script = ROOT / "scripts" / "plot_figure4_benchmark.py"
    subprocess.run([sys.executable, str(script)], cwd=ROOT, check=True)


def _format_table_auroc(auroc: float, ci: str) -> str:
    lo, hi = ci.split("-")
    return f"{float(auroc):.3f} ({lo}–{hi})"


def _load_tables() -> tuple[pd.DataFrame, pd.DataFrame]:
    for path in (FIGURE_PATH, TABLE2_CSV, TABLE3_CSV):
        if not path.exists():
            raise SystemExit(f"Missing artifact: {path}. Run plot_figure4_benchmark.py first.")
    t2 = pd.read_csv(TABLE2_CSV)
    t3 = pd.read_csv(TABLE3_CSV)
    return t2, t3


def _build_text_replacements(t2: pd.DataFrame, t3: pd.DataFrame) -> list[tuple[str, str]]:
    ev = t3.loc[t3["Model"] == "Ensemble_Voting"].iloc[0]
    lo, hi = str(ev["AUROC_95CI"]).split("-")
    auroc = f"{float(ev['AUROC']):.3f}"
    body_ci = f"{lo}-{hi}"
    table_ci = f"{lo}–{hi}"

    es = t2.loc[t2["Model"] == "Ensemble_Stacking"].iloc[0]
    es_lo, es_hi = str(es["AUROC_95CI"]).split("-")
    es_auroc = f"{float(es['AUROC']):.3f}"

    replacements = [
        ("0.722 (0.680-0.765)", f"{auroc} ({body_ci})"),
        (
            "an AUROC of 0.722 (95% CI 0.680-0.765; events=171)",
            f"an AUROC of {auroc} (95% CI {body_ci}; events=171)",
        ),
        (
            "achieved AUROC 0.722 (95% CI 0.680-0.765)",
            f"achieved AUROC {auroc} (95% CI {body_ci})",
        ),
        ("0.733 versus 0.722 for the 18-predictor", f"0.733 versus {auroc} for the 18-predictor"),
        (
            "Ensemble_Stacking also performed similarly with an AUROC of 0.841 (95% CI 0.768-0.914)",
            f"Ensemble_Stacking also performed similarly with an AUROC of {es_auroc} (95% CI {es_lo}-{es_hi})",
        ),
        ("an AUROC of 0.841 (95% CI 0.768-0.914)", f"an AUROC of {es_auroc} (95% CI {es_lo}-{es_hi})"),
        # Legacy embedded table text (no cell boundaries) — external rows that changed.
        ("Logistic Regression0.695 (0.652–0.738)", "Logistic Regression0.706 (0.662–0.749)"),
        ("ANN0.680 (0.631–0.729)", "ANN0.685 (0.636–0.735)"),
        ("XGBoost0.706 (0.662–0.749)", "XGBoost0.705 (0.662–0.749)"),
        ("GaussianNB0.662 (0.612–0.713)", "GaussianNB0.699 (0.653–0.746)"),
        (f"Ensemble_Voting0.722 ({table_ci})", f"Ensemble_Voting{auroc} ({table_ci})"),
        ("Ensemble_Stacking0.723 (0.681–0.765)", "Ensemble_Stacking0.727 (0.684–0.770)"),
    ]
    return replacements


def _replace_in_paragraph(paragraph, replacements: list[tuple[str, str]]) -> bool:
    text = paragraph.text
    new_text = text
    for old, new in replacements:
        if old in new_text:
            new_text = new_text.replace(old, new)
    if new_text != text:
        if paragraph.runs:
            paragraph.runs[0].text = new_text
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.text = new_text
        return True
    return False


def _patch_benchmark_table(table, df: pd.DataFrame) -> int:
    lookup = {row["Model"]: row for _, row in df.iterrows()}
    hits = 0
    for row in table.rows[1:]:
        model = row.cells[0].text.strip()
        if model not in lookup:
            continue
        src = lookup[model]
        new_val = _format_table_auroc(src["AUROC"], src["AUROC_95CI"])
        if row.cells[1].text.strip() != new_val:
            row.cells[1].text = new_val
            hits += 1
    return hits


def _find_benchmark_tables(doc: Document) -> tuple[object | None, object | None]:
    internal = external = None
    for table in doc.tables:
        flat = " ".join(c.text for row in table.rows for c in row.cells)
        if INTERNAL_MARKER in flat:
            internal = table
        elif EXTERNAL_MARKER in flat:
            external = table
    return internal, external


def _patch_docx_tables_and_text(path: Path, t2: pd.DataFrame, t3: pd.DataFrame) -> dict[str, int]:
    replacements = _build_text_replacements(t2, t3)
    doc = Document(path)
    stats = {"paragraphs": 0, "table_cells": 0, "benchmark_rows": 0}

    for paragraph in doc.paragraphs:
        if _replace_in_paragraph(paragraph, replacements):
            stats["paragraphs"] += 1

    internal, external = _find_benchmark_tables(doc)
    if internal is None or external is None:
        raise SystemExit(f"Could not locate Table 2/3 in {path.name}")

    stats["benchmark_rows"] += _patch_benchmark_table(internal, t2)
    stats["benchmark_rows"] += _patch_benchmark_table(external, t3)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if _replace_in_paragraph(paragraph, replacements):
                        stats["table_cells"] += 1

    doc.save(path)
    return stats


def _figure_to_tiff_bytes(png_path: Path) -> bytes:
    buf = io.BytesIO()
    Image.open(png_path).convert("RGB").save(buf, format="TIFF", compression="tiff_deflate")
    return buf.getvalue()


def _replace_docx_image(docx_path: Path, png_path: Path) -> None:
    tiff_bytes = _figure_to_tiff_bytes(png_path)
    tmp_path = docx_path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(tmp_path, "w") as zout:
        for info in zin.infolist():
            data = tiff_bytes if info.filename == MEDIA_IMAGE4 else zin.read(info.filename)
            zout.writestr(info, data)
    tmp_path.replace(docx_path)


def _verify_docx(path: Path, t3: pd.DataFrame) -> list[str]:
    errors: list[str] = []
    ev = f"{float(t3.loc[t3['Model'] == 'Ensemble_Voting', 'AUROC'].iloc[0]):.3f}"

    with zipfile.ZipFile(path) as z:
        if MEDIA_IMAGE4 not in z.namelist():
            errors.append(f"{path.name}: missing {MEDIA_IMAGE4}")
        xml = z.read("word/document.xml").decode("utf-8", errors="ignore")

    text = re.sub(r"<[^>]+>", "", xml)
    if "0.832" in text or "0.802" in text:
        errors.append(f"{path.name}: legacy Figure 4 text residue found")
    if "0.841" in text and "Ensemble_Stacking" in text:
        errors.append(f"{path.name}: stale Ensemble_Stacking 0.841 in body text")
    if "Ensemble_Voting0.722" in text.replace(" ", ""):
        errors.append(f"{path.name}: stale Ensemble_Voting 0.722 in embedded table text")
    if f"Ensemble_Voting{ev}" not in text.replace(" ", "") and f"Ensemble_Voting {ev}" not in text:
        errors.append(f"{path.name}: Ensemble_Voting {ev} not found in document text")

    doc = Document(path)
    _, external = _find_benchmark_tables(doc)
    if external is None:
        errors.append(f"{path.name}: external benchmark table missing after patch")
    else:
        for row in external.rows[1:]:
            if row.cells[0].text.strip() == "Ensemble_Voting":
                if not row.cells[1].text.startswith(ev):
                    errors.append(
                        f"{path.name}: Table 3 Ensemble_Voting is {row.cells[1].text}, expected {ev}"
                    )
    return errors


def main() -> None:
    parser = argparse.ArgumentParser(description="Patch low_fe_3 Figure 4 + Table 2/3 from FE1 benchmark")
    parser.add_argument("--regen-figure", action="store_true", help="Run plot_figure4_benchmark.py first")
    parser.add_argument("--backup", action="store_true", help="Save .bak copies before patching")
    args = parser.parse_args()

    if args.regen_figure:
        _run_benchmark_plot()

    t2, t3 = _load_tables()

    for manuscript in MANUSCRIPT_FILES:
        if not manuscript.exists():
            raise SystemExit(f"Missing manuscript: {manuscript}")
        if args.backup:
            shutil.copy2(manuscript, manuscript.with_suffix(".docx.bak"))

        stats = _patch_docx_tables_and_text(manuscript, t2, t3)
        _replace_docx_image(manuscript, FIGURE_PATH)
        errors = _verify_docx(manuscript, t3)

        print(f"Patched {manuscript.name}: {stats}")
        if errors:
            for err in errors:
                print(f"  VERIFY FAIL: {err}")
            raise SystemExit(1)
        print(f"  verify OK (Figure 4 image + Table 3 Ensemble_Voting synced)")


if __name__ == "__main__":
    main()
